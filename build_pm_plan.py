# -*- coding: utf-8 -*-
"""Rebuild data/pm_plan.json (and the PM part photos) from the PM Plan workbook.

The Planner's **Import templates** button does not read the .xlsx — it reads
data/pm_plan.json. This script is what turns one into the other.

    python build_pm_plan.py              # dry run: report + diff against the current json
    python build_pm_plan.py --apply      # write data/pm_plan.json and static/pmimg/

What it reads from every sheet (they are all laid out the same way):
    row 7      the header row: ลำดับ | รูปภาพ Part | รายการตรวจเช็ค | สถานะปกติ | Week … | วิธีการ
    row 9 →    one checklist item per row, until the ลำดับ legend row
    column A   the item number — and its FILL COLOUR is the frequency
    column B   the part photo, anchored on the item's row
    tab colour the template's colour group
Frequency legend (fill colour of the ลำดับ cell):
    FFFFC000 amber = weekly     FFFF0000 red   = monthly   FF00B050 green = 3-month
    FF00B0F0 blue  = 6-month    FFFFFF00 yellow / no fill  = every run
"""
import argparse, json, os, re, shutil, sys
from datetime import datetime

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "BFLFP PM Plan", "BFLFP PM Plan _R1.xlsx")
JSON = os.path.join(BASE, "data", "pm_plan.json")
IMGD = os.path.join(BASE, "static", "pmimg")
FACTORY_ID = 2

FREQ_BY_COLOR = {"FFFFC000": "weekly", "FFFF0000": "monthly", "FF00B050": "q3m",
                 "FF00B0F0": "m6", "FFFFFF00": "each_op"}
DEFAULT_FREQ = "each_op"          # no fill = the closing clean-up line


def _txt(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def _fill(cell):
    f = cell.fill
    if not f or not f.patternType:
        return ""
    c = f.fgColor
    return (c.rgb or "") if c is not None and c.type == "rgb" else ""


def _tab_color(ws):
    c = ws.sheet_properties.tabColor
    if c is None:
        return ""
    if c.type == "rgb":
        return c.rgb or ""
    if c.type == "theme":
        return "T%s" % c.theme          # a themed tab, kept as its own group
    return ""


def read_workbook(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)          # styles + images, so not read_only
    templates, images, form_no = [], {}, ""
    for idx, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        head = next((r for r in range(1, 16) if _txt(ws.cell(r, 1).value) == "ลำดับ"), None)
        if not head:
            continue                            # blank tab (Sheet1) — not a checklist
        form_no = form_no or _txt(ws.cell(3, 10).value)
        # the columns are not in the same place on every sheet — one has an extra
        # ความถี่ column, another shifts วิธีการ — so read them off the header row
        def col_of(label, default):
            return next((c for c in range(1, 22) if label in _txt(ws.cell(head, c).value)), default)
        icol = col_of("รายการตรวจเช็ค", 3)
        ncol = col_of("สถานะปกติ", 4)
        mcol = col_of("วิธีการ", 14)
        pcol = col_of("รูปภาพ", 2)
        stop = next((r for r in range(head + 1, (ws.max_row or head) + 1)
                     if "สัญลักษณ์" in _txt(ws.cell(r, 1).value)), (ws.max_row or head) + 1)
        # a photo belongs to the row it is anchored on; the one on the title rows is the logo
        by_row = {}
        for im in ws._images:
            r = im.anchor._from.row + 1
            if r > head:
                by_row.setdefault(r, im)
        items = []
        for r in range(head + 1, stop):
            if not _txt(ws.cell(r, 1).value).isdigit():
                continue                        # not an item row
            # the closing clean-up line writes its text in the photo column
            item = _txt(ws.cell(r, icol).value) or _txt(ws.cell(r, pcol).value)
            if not item:
                continue                        # a spare numbered row with nothing on it
            # numbered fresh: several sheets repeat or skip a number by hand
            seq = len(items) + 1
            img = ""
            if r in by_row:
                ext = (getattr(by_row[r], "path", "") or ".jpg").rsplit(".", 1)[-1].lower()
                ext = "jpg" if ext in ("jpeg", "jpg") else ext
                img = "t%02d_%d.%s" % (idx, seq, ext)
                images[img] = by_row[r]
            items.append({"seq": seq, "item": item,
                          "normal": _txt(ws.cell(r, ncol).value),
                          "method": _txt(ws.cell(r, mcol).value),
                          "freq": FREQ_BY_COLOR.get(_fill(ws.cell(r, 1)), DEFAULT_FREQ),
                          "img": img})
        if items:
            templates.append({"machine_type": name.strip(), "color": _tab_color(ws), "items": items})
    return templates, images, form_no


def color_groups(templates):
    from collections import Counter
    n = Counter(t["color"] for t in templates if t["color"])
    return [{"color": c, "name": "Group %d" % (i + 1), "n": n[c]}
            for i, (c, _) in enumerate(n.most_common())]


def diff_against_current(templates):
    if not os.path.exists(JSON):
        print("  (no current pm_plan.json to compare with)")
        return
    old = json.load(open(JSON, encoding="utf-8"))
    o = {t["machine_type"]: t for t in old["templates"]}
    n = {t["machine_type"]: t for t in templates}
    added, gone = sorted(set(n) - set(o)), sorted(set(o) - set(n))
    print("  sheets added : %d  %s" % (len(added), ", ".join(added) or "—"))
    print("  sheets gone  : %d  %s" % (len(gone), ", ".join(gone) or "—"))
    changed = 0
    for name in sorted(set(o) & set(n)):
        a = [(i["seq"], i["item"], i["normal"], i["method"], i["freq"]) for i in o[name]["items"]]
        b = [(i["seq"], i["item"], i["normal"], i["method"], i["freq"]) for i in n[name]["items"]]
        if a != b:
            changed += 1
            print("  changed: %s" % name)
            for x, y in zip(a, b):
                if x != y:
                    print("      was %s\n      now %s" % (x, y))
            if len(a) != len(b):
                print("      %d items -> %d items" % (len(a), len(b)))
    print("  sheets identical to the current json: %d of %d" % (len(set(o) & set(n)) - changed, len(set(o) & set(n))))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write the files (default: dry run)")
    ap.add_argument("--xlsx", default=XLSX)
    a = ap.parse_args()
    if not os.path.exists(a.xlsx):
        sys.exit("workbook not found: %s" % a.xlsx)

    templates, images, form_no = read_workbook(a.xlsx)
    items = sum(len(t["items"]) for t in templates)
    from collections import Counter
    freqs = Counter(i["freq"] for t in templates for i in t["items"])
    print("read %s" % os.path.basename(a.xlsx))
    print("  templates %d | items %d | photos %d | form %s" % (len(templates), items, len(images), form_no))
    print("  frequencies: %s" % dict(freqs))
    print("  colour groups: %s" % [(g["color"], g["n"]) for g in color_groups(templates)])
    print("\ncompared with the current data/pm_plan.json:")
    diff_against_current(templates)

    plan = {"form_no": form_no, "factory_id": FACTORY_ID,
            "templates": templates, "color_groups": color_groups(templates)}
    if not a.apply:
        print("\nDRY RUN — nothing written. Re-run with --apply to write:")
        print("  %s" % JSON)
        print("  %s  (%d photos)" % (IMGD, len(images)))
        return
    if os.path.exists(JSON):
        bak = JSON + ".bak-" + datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(JSON, bak)
        print("\nbacked up the old json -> %s" % os.path.basename(bak))
    os.makedirs(IMGD, exist_ok=True)
    for fname, im in images.items():
        with open(os.path.join(IMGD, fname), "wb") as fh:
            fh.write(im._data())
    json.dump(plan, open(JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("wrote %s and %d photos into %s" % (JSON, len(images), IMGD))
    print("\nNow press Planner -> Import templates to load it.")


if __name__ == "__main__":
    main()
