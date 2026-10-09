"""Generate printable QR label sheets for every asset — one PDF per factory.

Run on the PC (Windows), so the Thai machine names render:
    pip install reportlab openpyxl
    python make_qr_labels.py            # reads the 3 Excel from ./assets
    python make_qr_labels.py C:\\path   # or another folder holding them

Each label = a QR that encodes the asset code, plus the code and the machine name.
Output: QR-Labels-BFL.pdf / QR-Labels-BFLFP.pdf / QR-Labels-BFLPC.pdf
"""
import os, sys, glob

try:
    import openpyxl
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as canvaslib
    from reportlab.graphics.barcode import qr
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics import renderPDF
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import simpleSplit
except ImportError:
    sys.exit("Missing libraries — run:  pip install reportlab openpyxl")

FILE_TITLE = {"BFL": "BFL Main Plant", "BFLFP": "FP Wet Food", "BFLPC": "PC Petcare"}
COLS, ROWS = 3, 8

THF = "Helvetica"
for p in [r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\LeelawUI.ttf",
          r"C:\Windows\Fonts\leelawad.ttf",
          "/usr/share/fonts/truetype/tlwg/Garuda.ttf"]:
    if os.path.exists(p):
        pdfmetrics.registerFont(TTFont("TH", p)); THF = "TH"; break


def gen(items, out, title):
    c = canvaslib.Canvas(out, pagesize=A4); W, H = A4; m = 8 * mm
    cw = (W - 2 * m) / COLS; ch = (H - 2 * m - 8 * mm) / ROWS; per = COLS * ROWS
    for i, (code, name) in enumerate(items):
        p = i % per
        if p == 0:
            if i > 0:
                c.showPage()
            c.setFont("Helvetica-Bold", 11)
            c.drawString(m, H - m + 1 * mm, f"{title} - {len(items)} asset QR labels")
        col = p % COLS; row = p // COLS
        x = m + col * cw; y = H - m - 8 * mm - (row + 1) * ch
        c.setStrokeColorRGB(.82, .82, .82); c.rect(x + 1.5 * mm, y + 1.5 * mm, cw - 3 * mm, ch - 3 * mm)
        q = min(ch - 6 * mm, 24 * mm)
        w = qr.QrCodeWidget(code); b = w.getBounds(); bw = b[2] - b[0]; bh = b[3] - b[1]
        d = Drawing(q, q, transform=[q / bw, 0, 0, q / bh, 0, 0]); d.add(w)
        renderPDF.draw(d, c, x + 3 * mm, y + (ch - q) / 2)
        tx = x + q + 6 * mm; tw = cw - q - 9 * mm
        c.setFont("Helvetica-Bold", 11); c.drawString(tx, y + ch - 8 * mm, code)
        c.setFont(THF, 7)
        for j, ln in enumerate(simpleSplit(name, THF, 7, tw)[:3]):
            c.drawString(tx, y + ch - 13 * mm - j * 3.4 * mm, ln)
    c.showPage(); c.save()


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
    files = sorted(glob.glob(os.path.join(folder, "Asset-Review-*.xlsx")))
    if not files:
        sys.exit(f"No Asset-Review-*.xlsx in {folder} — put the 3 files in an 'assets' folder.")
    if THF == "Helvetica":
        print("! No Thai font found — machine names may not render (expected Windows Tahoma/Leelawadee).")
    for f in files:
        suf = os.path.basename(f)[len("Asset-Review-"):-len(".xlsx")]
        title = FILE_TITLE.get(suf, suf)
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True); ws = wb["Review"]
        items = [(str(r[1]).strip(), str(r[2]).strip() if r[2] else "")
                 for r in ws.iter_rows(min_row=6, values_only=True)
                 if r[1] and str(r[1]).strip() != "W01XX99"]
        wb.close()
        out = f"QR-Labels-{suf}.pdf"
        gen(items, out, title)
        print(f"  {out}: {len(items)} labels")


if __name__ == "__main__":
    main()
