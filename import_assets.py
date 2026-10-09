"""Import the factory asset registers into the machines (assets) table.

Run once from the project root:
    pip install openpyxl
    python import_assets.py            # reads the 3 files from ./assets
    python import_assets.py C:\\path   # or point it at another folder

Expects these files (as reviewed):
    Asset-Review-BFL.xlsx    -> factory BFL
    Asset-Review-BFLFP.xlsx  -> factory FP
    Asset-Review-BFLPC.xlsx  -> factory PC

What it does:
  * adds the extra register columns to `machines` if they are missing
  * reads each sheet 'Review' from row 6 (skipping title/header/example rows)
  * UPSERTS by asset code — safe to re-run; it updates existing rows and never
    creates duplicates, so job links to a machine stay intact
  * prints a verification report you can check against the numbers below.

Expected result: BFL 716 · FP 151 · PC 731 · total 1598 assets.
"""
import os, sys, glob, sqlite3, datetime
from server.config import DB_PATH

FILE_FACTORY = {"BFL": "BFL", "BFLFP": "FP", "BFLPC": "PC"}   # filename suffix -> factory code
NEW_COLS = ["category", "asset_group", "floor", "department", "manufacturer", "size", "remark"]
# machines field  <-  Excel column (0-based index; B=1, C=2, ...)
MAP = [("code", 1), ("name", 2), ("category", 3), ("asset_group", 4), ("kpi_class", 5),
       ("criticality", 6), ("line", 7), ("floor", 8), ("department", 9), ("brand_model", 10),
       ("manufacturer", 11), ("size", 12), ("serial_no", 13), ("year_install", 14),
       ("last_pm_date", 15), ("remark", 16)]
FIELDS = [m[0] for m in MAP] + ["factory_id", "active"]


def cell(v):
    if v in (None, ""):
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def main():
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl not installed — run:  pip install openpyxl")

    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
    files = sorted(glob.glob(os.path.join(folder, "Asset-Review-*.xlsx")))
    if not files:
        sys.exit(f"No Asset-Review-*.xlsx found in: {folder}\n"
                 f"Put the 3 files in an 'assets' folder, or pass the folder path.")

    db = sqlite3.connect(DB_PATH)
    c = db.cursor()

    have = {r[1] for r in c.execute("PRAGMA table_info(machines)")}
    for col in NEW_COLS:
        if col not in have:
            c.execute(f"ALTER TABLE machines ADD COLUMN {col} TEXT DEFAULT ''")

    fac_id = {code: fid for fid, code in c.execute("SELECT id, code FROM factories")}

    placeholders = ",".join("?" * len(FIELDS))
    setclause = ",".join(f"{f}=excluded.{f}" for f in FIELDS)
    sql = (f"INSERT INTO machines({','.join(FIELDS)}) VALUES({placeholders}) "
           f"ON CONFLICT(code) DO UPDATE SET {setclause}")

    inserted = updated = skipped = 0
    per_file = {}
    for f in files:
        suffix = os.path.basename(f)[len("Asset-Review-"):-len(".xlsx")]
        fcode = FILE_FACTORY.get(suffix)
        if fcode not in fac_id:
            print(f"  ! skipping {os.path.basename(f)} — unknown factory '{suffix}'")
            continue
        fid = fac_id[fcode]
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb["Review"]
        n = 0
        for row in ws.iter_rows(min_row=6, values_only=True):
            code = cell(row[1])
            if not code or code == "W01XX99":          # empty / example row
                skipped += 1
                continue
            rec = {field: cell(row[idx]) for field, idx in MAP}
            rec["factory_id"] = fid
            rec["active"] = 1
            exists = c.execute("SELECT 1 FROM machines WHERE code=?", (code,)).fetchone()
            c.execute(sql, [rec[f] for f in FIELDS])
            if exists:
                updated += 1
            else:
                inserted += 1
            n += 1
        wb.close()
        per_file[os.path.basename(f)] = (fcode, n)

    db.commit()

    print("\n=== import complete ===")
    for name, (fcode, n) in per_file.items():
        print(f"  {name:28s} {fcode:5s} {n} rows read")
    print(f"  inserted (new): {inserted}   updated (existing): {updated}   skipped: {skipped}")
    print("\n=== verification (machines table now) ===")
    fname = {v: k for k, v in fac_id.items()}
    for fid, code in c.execute("SELECT id, code FROM factories ORDER BY id"):
        n = c.execute("SELECT COUNT(*) FROM machines WHERE factory_id=?", (fid,)).fetchone()[0]
        print(f"  factory {code:5s} (id {fid}): {n} assets")
    print("  total assets:", c.execute("SELECT COUNT(*) FROM machines").fetchone()[0])
    print("  by KPI class:", dict(c.execute("SELECT kpi_class, COUNT(*) FROM machines GROUP BY kpi_class").fetchall()))
    print("  by criticality:", dict(c.execute("SELECT criticality, COUNT(*) FROM machines GROUP BY criticality").fetchall()))
    print("  distinct codes:", c.execute("SELECT COUNT(DISTINCT code) FROM machines").fetchone()[0])
    print("\nExpected from the registers: BFL 716 · FP 151 · PC 731 · total 1598 (plus any older demo machines).")
    db.close()


if __name__ == "__main__":
    main()
