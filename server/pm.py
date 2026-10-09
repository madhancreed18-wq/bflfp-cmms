"""Preventive-Maintenance.

Three layers:
  1. pm_templates + pm_items  — the checklist LIBRARY imported from the BFLFP PM
     Plan workbook. Each item carries its frequency (weekly / monthly / 3-month /
     6-month / yearly / every-run), read from the colour coding on the form.
  2. pm_plans + pm_plan_sched — the PLAN: the planner assigns a template to one
     machine, sets a start date, and confirms an interval per frequency.
  3. Generation — a daily server task turns due plan schedules into PM jobs, which
     flow through the normal Plan -> assign -> sign-off workflow.
"""
import os, json, re, math
from contextlib import closing
from datetime import datetime, date, timedelta

from fastapi import APIRouter, Request, HTTPException

from .db import (db, now, today, log_status, job_row, next_jobid, set_stage,
                 user_names, VOID_SQL, offdays_of)
from .auth import require_role

import logging
_log = logging.getLogger("cmms")

router = APIRouter(prefix="/api/pm")

# A PM job in one of these has been started (or finished): its checklist is frozen.
STARTED = ("InProgress", "Paused", "Hold", "Rework", "ServiceCompleted", "Done",
           "Rejected", "Cancelled")
_STARTED_SQL = "(" + ",".join("'%s'" % x for x in STARTED) + ")"
# the live sheet only ever shows items still on it
ACTIVE_ITEM = " AND COALESCE(active,1)=1"

PLAN_FILE = os.path.join(os.path.dirname(__file__), "..", "data", "pm_plan.json")

FREQ_DAYS = {"weekly": 7, "monthly": 30, "q3m": 90, "m6": 180, "yearly": 365, "each_op": 0}
FREQ_ORDER = {"weekly": 0, "monthly": 1, "q3m": 2, "m6": 3, "yearly": 4, "each_op": 5}
FREQ_LABEL = {"weekly": "Weekly", "monthly": "Monthly", "q3m": "3-month",
              "m6": "6-month", "yearly": "Yearly", "each_op": "Every run"}


_RECORDS_READY = False


def _ensure_records(c):
    """Tables that keep a PM record whole whatever later happens to the checklist.

    pm_results used to hold only an item id. An Excel upload deleted and recreated every
    item, so after one re-upload a finished PM showed its checklist as unanswered — on
    screen and on the PDF. Now:
      * every answer also stores the wording, normal state, method and frequency it was
        answered against (item_text …);
      * pm_job_items freezes the list a PM job STARTED with, so an edit made while a
        technician is mid-round never adds a row that blocks his Stop button;
      * pm_item_log is the change history shown on the sheet.
    The one-off fill below runs once per server start and only touches rows still empty.
    """
    global _RECORDS_READY
    if _RECORDS_READY:
        return
    for col, ddl in (("item_text", "TEXT DEFAULT ''"), ("item_normal", "TEXT DEFAULT ''"),
                     ("item_method", "TEXT DEFAULT ''"), ("item_freq", "TEXT DEFAULT ''")):
        try:
            c.execute(f"SELECT {col} FROM pm_results LIMIT 1")
        except Exception:
            try:
                c.execute(f"ALTER TABLE pm_results ADD COLUMN {col} {ddl}")
            except Exception:
                pass
    c.execute("""CREATE TABLE IF NOT EXISTS pm_job_items(
        id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER, item_id INTEGER, seq INTEGER,
        item TEXT DEFAULT '', normal_status TEXT DEFAULT '', method TEXT DEFAULT '',
        img TEXT DEFAULT '', freq TEXT DEFAULT '')""")
    try:
        c.execute("CREATE INDEX IF NOT EXISTS ix_pm_job_items_job ON pm_job_items(job_id)")
    except Exception:
        pass
    c.execute("""CREATE TABLE IF NOT EXISTS pm_item_log(
        id INTEGER PRIMARY KEY AUTOINCREMENT, factory_id INTEGER, template_id INTEGER,
        item_id INTEGER, action TEXT, before TEXT DEFAULT '', after TEXT DEFAULT '',
        user_id INTEGER, user_name TEXT DEFAULT '', at TEXT DEFAULT '')""")
    try:
        # answers saved before this existed: copy in the wording while the items still exist
        c.execute("""UPDATE pm_results SET
              item_text  =(SELECT i.item FROM pm_items i WHERE i.id=pm_results.item_id),
              item_normal=(SELECT COALESCE(i.normal_status,'') FROM pm_items i WHERE i.id=pm_results.item_id),
              item_method=(SELECT COALESCE(i.method,'') FROM pm_items i WHERE i.id=pm_results.item_id),
              item_freq  =(SELECT COALESCE(i.freq,'') FROM pm_items i WHERE i.id=pm_results.item_id)
            WHERE COALESCE(item_text,'')='' AND item_id IN (SELECT id FROM pm_items)""")
        c.commit()
    except Exception as e:
        _log.warning("[pm] answer wording back-fill skipped: %s", e)
    _RECORDS_READY = True
    try:
        # PM jobs already started or finished: freeze the list they were worked from
        n = 0
        for j in c.execute("SELECT j.* FROM jobs j WHERE j.jobtype='PM' AND j.machine_id IS NOT NULL"
                           " AND (j.status IN " + _STARTED_SQL +
                           "      OR j.id IN (SELECT job_id FROM pm_results))"
                           " AND j.id NOT IN (SELECT job_id FROM pm_job_items)").fetchall():
            if snapshot_job_items(c, dict(j)):
                n += 1
        c.commit()
        if n:
            _log.info("[pm] checklist frozen on %d started/finished PM job(s)", n)
    except Exception as e:
        _log.warning("[pm] job checklist freeze skipped: %s", e)


def _ensure(c):
    c.execute("""CREATE TABLE IF NOT EXISTS pm_templates(
        id INTEGER PRIMARY KEY AUTOINCREMENT, factory_id INTEGER DEFAULT 2,
        machine_type TEXT, form_no TEXT DEFAULT '')""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_items(
        id INTEGER PRIMARY KEY AUTOINCREMENT, template_id INTEGER, seq INTEGER,
        item TEXT, normal_status TEXT, method TEXT, freq TEXT DEFAULT 'weekly')""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_plans(
        id INTEGER PRIMARY KEY AUTOINCREMENT, factory_id INTEGER DEFAULT 2,
        template_id INTEGER, machine_id INTEGER, active INTEGER DEFAULT 1,
        start_date TEXT DEFAULT '', created_by INTEGER, created_at TEXT DEFAULT '')""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_plan_sched(
        id INTEGER PRIMARY KEY AUTOINCREMENT, plan_id INTEGER, freq TEXT,
        interval_days INTEGER DEFAULT 7, active INTEGER DEFAULT 1, last_gen TEXT DEFAULT '')""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_members(
        id INTEGER PRIMARY KEY AUTOINCREMENT, template_id INTEGER, machine_id INTEGER, factory_id INTEGER DEFAULT 2)""")
    # What a technician actually found, one row per checklist item on one PM job.
    # The sheet is the record: OK or NG, why if NG, and whose hands were on it —
    # booked to the person named on the phone, not to the shared login.
    c.execute("""CREATE TABLE IF NOT EXISTS pm_results(
        id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER, item_id INTEGER,
        seq INTEGER DEFAULT 0, result TEXT DEFAULT '', remark TEXT DEFAULT '',
        tech INTEGER, tech_name TEXT DEFAULT '', at TEXT DEFAULT '')""")
    try:
        c.execute("CREATE UNIQUE INDEX IF NOT EXISTS ix_pm_results_job_item"
                  " ON pm_results(job_id,item_id)")
    except Exception:
        pass
    c.execute("""CREATE TABLE IF NOT EXISTS pm_marks(
        id INTEGER PRIMARY KEY AUTOINCREMENT, machine_id INTEGER, item_id INTEGER, cell TEXT)""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_config(
        factory_id INTEGER PRIMARY KEY, start_date TEXT DEFAULT '', active INTEGER DEFAULT 0)""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_overrides(
        id INTEGER PRIMARY KEY AUTOINCREMENT, machine_id INTEGER, orig_date TEXT, new_date TEXT)""")
    c.execute("""CREATE TABLE IF NOT EXISTS pm_groups(
        id INTEGER PRIMARY KEY AUTOINCREMENT, factory_id INTEGER, color TEXT, name TEXT)""")
    try:
        c.execute("SELECT color FROM pm_templates LIMIT 1")
    except Exception:
        try:
            c.execute("ALTER TABLE pm_templates ADD COLUMN color TEXT DEFAULT ''")
        except Exception:
            pass
    for col, ddl in (("freq", "ALTER TABLE pm_items ADD COLUMN freq TEXT DEFAULT 'weekly'"),
                     ("img", "ALTER TABLE pm_items ADD COLUMN img TEXT DEFAULT ''"),
                     # an item taken off the sheet is hidden, never deleted: past answers
                     # and started jobs still point at it
                     ("active", "ALTER TABLE pm_items ADD COLUMN active INTEGER DEFAULT 1"),
                     # changed in the app since the last Excel upload — the upload warns
                     # before it writes the file's wording over it
                     ("app_edit", "ALTER TABLE pm_items ADD COLUMN app_edit INTEGER DEFAULT 0")):
        try:
            c.execute(f"SELECT {col} FROM pm_items LIMIT 1")
        except Exception:
            try:
                c.execute(ddl)
            except Exception:
                pass
    _ensure_records(c)
    # b413: electrical points (yellow in the master) and the CE team's part of a job
    try:
        from .elec import ensure as _el_ensure
        _el_ensure(c)
    except Exception as e:
        _log.warning("[pm] electrical columns: %s", e)


    # The Assets grid lets a planner pin an asset's PM group by hand. Empty means
    # "follow the colour of the PM template this asset matches" — the old behaviour.
    try:
        c.execute("SELECT pm_group_color FROM machines LIMIT 1")
    except Exception:
        try:
            c.execute("ALTER TABLE machines ADD COLUMN pm_group_color TEXT DEFAULT ''")
            c.commit()    # a read-only request must not sit on the write lock
        except Exception:
            pass


def _fac(u):
    return u.get("active_factory") or u.get("factory_id")


def _toks(s):
    return [w for w in re.split(r"[\s/()#0-9.\-]+", s or "") if len(w) >= 3]


def _norm(s):
    """Collapse a name to Thai+alnum only, drop trailing serial (#1) — for robust
    asset-name <-> sheet-name matching."""
    s = re.sub(r"[^0-9a-z฀-๿]", "", (s or "").lower())
    return re.sub(r"\d+$", "", s)


# Spelling variants: (substring found in the asset name) -> (the sheet it means).
# The register and the PM sheets were written by different hands, so the same machine
# is spelled two ways. Keep the key specific enough that it cannot fire on another
# machine — "ปั๊มวัติถุดิบ" not "วัติถุดิบ", which also appears in a room name.
PM_ALIASES = [("damper", "Dumper"), ("ยิงโค", "เครื่องยิงโค้ด"), ("sealถุง", "เครื่องซีลถุง"),
              ("นึ่งเนื้อไก่", "ตู้นึ่งไก่"), ("เอียงตระกร้า", "เครื่องเอียงตะกร้า"),
              ("คอยเย็นห้อง", "คอยเย็นแอร์"),              # 11 fan coils named by their room
              ("ปั๊มวัติถุดิบ", "ปั๊มวัตถุดิบ"),            # register misspells วัตถุดิบ
              ("เครื่งตรวจจับโลหะ", "เครื่องตรวจจับโลหะแบบสายพาน"),   # register drops the อ
              ("สเปร์คูลลิ่ง", "รางสเปร์ยคูลลิ่งปลา"),
              ("สายพานรับซอง", "สายพานรับซองเพ้าช์"),
              ("ฆ่าเชื้อกล่อง", "หม้อฆ่าเชื้อกล่อง"),
              ("นึ่งแรงดันไอน้ำ", "หม้อฆ่าเชื้อเพ๊าท์"),    # by elimination: ST02=กระป๋อง, ST05=กล่อง
              ("แมวเลียttk", "บรรจุแมวเลียTTK"),           # family: the number picks TTK-1 / TTK-2
              ("แมวเลียtony", "บรรจุแมวเลีย Tony")]


def _alias_bonus(machine_type, an):
    """Matched on the normalised name, so an alias can point at a family of numbered
    sheets (บรรจุแมวเลียTTK-1 / -2) and let the machine number choose between them."""
    for k, tmt in PM_ALIASES:
        if _norm(machine_type) == _norm(tmt) and _norm(k) in an:
            return 1500
    return 0


def _tail_no(s):
    """The machine number in a name — #3, -3, or a trailing 3. Used only to choose
    between sheets that are otherwise identically named."""
    m = re.search(r"#\s*(\d+)", s or "")
    if m:
        return int(m.group(1))
    m = re.search(r"(\d+)\s*$", (s or "").strip())
    return int(m.group(1)) if m else None


def _next_due(interval_days, start_date, last_gen):
    iv = int(interval_days or 0)
    if iv <= 0:
        return None
    if last_gen:
        try:
            return (date.fromisoformat(last_gen[:10]) + timedelta(days=iv)).isoformat()
        except Exception:
            return today()
    return (start_date or today())[:10]


# ---------- library ----------

@router.get("/templates")
async def list_templates(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT id,machine_type,
              (SELECT COUNT(*) FROM pm_items WHERE template_id=t.id AND COALESCE(active,1)=1) n_items,
              (SELECT COUNT(*) FROM pm_plans WHERE template_id=t.id AND active=1) n_plans
            FROM pm_templates t WHERE factory_id=? ORDER BY machine_type""", (fac,)).fetchall()
        out = []
        for r in rows:
            fq = c.execute("SELECT freq,COUNT(*) n FROM pm_items WHERE template_id=?" + ACTIVE_ITEM + " GROUP BY freq", (r["id"],)).fetchall()
            fl = [{"freq": x["freq"], "label": FREQ_LABEL.get(x["freq"], x["freq"]), "n": x["n"]} for x in fq]
            fl.sort(key=lambda x: FREQ_ORDER.get(x["freq"], 9))
            out.append({**dict(r), "freqs": fl})
    return {"templates": out, "total": len(out)}


@router.get("/templates/{tid}")
async def template_detail(tid: int, req: Request):
    require_role(req, "planner", "admin", "manager", "technician")
    with closing(db()) as c:
        _ensure(c)
        t = c.execute("SELECT * FROM pm_templates WHERE id=?", (tid,)).fetchone()
        if not t:
            raise HTTPException(404, "template not found")
        groups = {}
        for it in c.execute("SELECT seq,item,normal_status,method,freq,img FROM pm_items WHERE template_id=?" + ACTIVE_ITEM + " ORDER BY seq", (tid,)):
            groups.setdefault(it["freq"], []).append(dict(it))
        g = [{"freq": f, "label": FREQ_LABEL.get(f, f), "default_days": FREQ_DAYS.get(f, 30), "items": v}
             for f, v in groups.items()]
        g.sort(key=lambda x: FREQ_ORDER.get(x["freq"], 9))
    return {"id": t["id"], "machine_type": t["machine_type"], "groups": g}


@router.get("/machines")
async def pm_machines(req: Request):
    """Machines for the assign picker. With no search text, the list is filtered to
    the machines that match the chosen checklist's type (e.g. COOLING WATER ->
    only the cooling-water assets). A search overrides and looks across all assets."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    q = (req.query_params.get("q") or "").strip()
    tid = req.query_params.get("template_id")
    with closing(db()) as c:
        _ensure(c)
        mtype = ""
        if tid:
            r = c.execute("SELECT machine_type FROM pm_templates WHERE id=?", (tid,)).fetchone()
            mtype = r["machine_type"] if r else ""
        allrows = [dict(r) for r in c.execute(
            "SELECT id,code,name,asset_group,category FROM machines WHERE factory_id=? ORDER BY code", (fac,))]
        matched = False
        if q:
            ql = q.lower()
            rows = [r for r in allrows if ql in (r["code"] or "").lower() or ql in (r["name"] or "").lower()]
        elif mtype:
            toks = [t.lower() for t in _toks(mtype)]

            def hay(r):
                return " ".join([r["name"] or "", r["asset_group"] or "", r["category"] or "", r["code"] or ""]).lower()
            rows = [r for r in allrows if toks and any(t in hay(r) for t in toks)]
            matched = bool(rows)
            if not rows:
                rows = allrows                       # no type match -> let them pick from all
        else:
            rows = allrows
        rows = rows[:400]
        planned = set()
        if tid:
            planned = {r["machine_id"] for r in c.execute(
                "SELECT machine_id FROM pm_plans WHERE template_id=? AND active=1", (tid,))}
        out = [{"id": r["id"], "code": r["code"], "name": r["name"], "planned": r["id"] in planned} for r in rows]
    return {"machines": out, "machine_type": mtype, "matched": matched, "count": len(out)}


# ---------- plans ----------

@router.post("/plans")
async def create_plan(req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    tid = b.get("template_id"); mid = b.get("machine_id")
    if not tid or not mid:
        raise HTTPException(400, "template_id and machine_id required")
    start = (b.get("start_date") or today())[:10]
    scheds = b.get("scheds") or []
    with closing(db()) as c:
        _ensure(c)
        if not c.execute("SELECT 1 FROM machines WHERE id=? AND factory_id=?", (mid, fac)).fetchone():
            raise HTTPException(400, "machine not in this factory")
        ex = c.execute("SELECT id FROM pm_plans WHERE template_id=? AND machine_id=?", (tid, mid)).fetchone()
        if ex:
            pid = ex["id"]
            c.execute("UPDATE pm_plans SET active=1,start_date=? WHERE id=?", (start, pid))
            c.execute("DELETE FROM pm_plan_sched WHERE plan_id=?", (pid,))
        else:
            pid = c.insert_id("INSERT INTO pm_plans(factory_id,template_id,machine_id,active,start_date,created_by,created_at)"
                              " VALUES(?,?,?,1,?,?,?)", (fac, tid, mid, start, u["id"], now()))
        present = {r["freq"] for r in c.execute("SELECT DISTINCT freq FROM pm_items WHERE template_id=?" + ACTIVE_ITEM, (tid,))}
        for s in scheds:
            fq = s.get("freq")
            if fq not in present:
                continue
            c.execute("INSERT INTO pm_plan_sched(plan_id,freq,interval_days,active,last_gen) VALUES(?,?,?,?,'')",
                      (pid, fq, max(0, int(s.get("interval_days") or FREQ_DAYS.get(fq, 30))),
                       1 if s.get("active", True) else 0))
        c.commit()
    return {"ok": True, "plan_id": pid}


@router.get("/plans")
async def list_plans(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    td = today()
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT p.id,p.template_id,p.machine_id,p.active,p.start_date,
              t.machine_type, m.code mcode, m.name mname
            FROM pm_plans p JOIN pm_templates t ON t.id=p.template_id
            LEFT JOIN machines m ON m.id=p.machine_id
            WHERE p.factory_id=? ORDER BY m.code""", (fac,)).fetchall()
        out = []
        for r in rows:
            sc = []
            for s in c.execute("SELECT * FROM pm_plan_sched WHERE plan_id=?", (r["id"],)):
                nd = _next_due(s["interval_days"], r["start_date"], s["last_gen"])
                sc.append({"id": s["id"], "freq": s["freq"], "label": FREQ_LABEL.get(s["freq"], s["freq"]),
                           "interval_days": s["interval_days"], "active": s["active"],
                           "last_gen": s["last_gen"], "next_due": nd,
                           "due": bool(nd) and bool(s["active"]) and nd <= td})
            sc.sort(key=lambda x: FREQ_ORDER.get(x["freq"], 9))
            out.append({**dict(r), "scheds": sc, "due_count": sum(1 for x in sc if x["due"])})
    return {"plans": out}


@router.get("/overview")
async def overview(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    td = date.fromisoformat(today())
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT ps.interval_days,ps.last_gen,ps.active,p.start_date,p.machine_id
            FROM pm_plan_sched ps JOIN pm_plans p ON p.id=ps.plan_id
            WHERE p.factory_id=? AND p.active=1""", (fac,)).fetchall()
    buckets = {"today": 0, "week": 0, "month": 0, "later": 0}
    m_today, m_week = set(), set()
    plans = 0
    for r in rows:
        if not r["active"]:
            continue
        nd = _next_due(r["interval_days"], r["start_date"], r["last_gen"])
        if not nd:
            continue
        plans += 1
        d = date.fromisoformat(nd)
        days = (d - td).days
        if days <= 0:
            buckets["today"] += 1; m_today.add(r["machine_id"]); m_week.add(r["machine_id"])
        elif days <= 7:
            buckets["week"] += 1; m_week.add(r["machine_id"])
        elif days <= 31:
            buckets["month"] += 1
        else:
            buckets["later"] += 1
    machines_planned = len({r["machine_id"] for r in rows})
    return {"due_today": buckets["today"], "due_week": buckets["today"] + buckets["week"],
            "due_month": buckets["today"] + buckets["week"] + buckets["month"],
            "machines_due_today": len(m_today), "machines_due_week": len(m_week),
            "scheduled_plans": plans, "machines_planned": machines_planned}


@router.get("/machine-tree")
async def machine_tree(req: Request):
    """Machines grouped by checklist group (membership), plus an ungrouped bucket —
    for the tree-view machine list."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        machines = {r["id"]: {"id": r["id"], "code": r["code"], "name": r["name"]}
                    for r in c.execute("SELECT id,code,name FROM machines WHERE factory_id=? ORDER BY code", (fac,))}
        groups, grouped = {}, set()
        for r in c.execute("""SELECT pm.template_id,pm.machine_id,t.machine_type
                FROM pm_members pm JOIN pm_templates t ON t.id=pm.template_id
                WHERE pm.factory_id=? ORDER BY t.machine_type""", (fac,)):
            if r["machine_id"] not in machines:
                continue
            g = groups.setdefault(r["template_id"], {"id": r["template_id"], "machine_type": r["machine_type"], "machines": []})
            g["machines"].append(machines[r["machine_id"]])
            grouped.add(r["machine_id"])
        ungrouped = [m for mid, m in machines.items() if mid not in grouped]
    return {"groups": sorted(groups.values(), key=lambda x: x["machine_type"]),
            "ungrouped": ungrouped, "total": len(machines)}


@router.get("/groups")
async def groups(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT id,machine_type,
              (SELECT COUNT(*) FROM pm_items WHERE template_id=t.id AND COALESCE(active,1)=1) n_items,
              (SELECT COUNT(*) FROM pm_members WHERE template_id=t.id) n_members
            FROM pm_templates t WHERE factory_id=? ORDER BY machine_type""", (fac,)).fetchall()
    return {"groups": [dict(r) for r in rows]}


@router.get("/groups/{tid}")
async def group_detail(tid: int, req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        t = c.execute("SELECT id,machine_type FROM pm_templates WHERE id=? AND factory_id=?", (tid, fac)).fetchone()
        if not t:
            raise HTTPException(404, "group not found")
        members = [dict(r) for r in c.execute(
            "SELECT m.id,m.code,m.name FROM pm_members pm JOIN machines m ON m.id=pm.machine_id "
            "WHERE pm.template_id=? ORDER BY m.code", (tid,))]
    return {"id": t["id"], "machine_type": t["machine_type"], "members": members}


@router.post("/groups/{tid}/members")
async def add_member(tid: int, req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    mid = b.get("machine_id")
    with closing(db()) as c:
        _ensure(c)
        if not c.execute("SELECT 1 FROM machines WHERE id=? AND factory_id=?", (mid, fac)).fetchone():
            raise HTTPException(400, "machine not in this factory")
        if not c.execute("SELECT 1 FROM pm_members WHERE template_id=? AND machine_id=?", (tid, mid)).fetchone():
            c.execute("INSERT INTO pm_members(template_id,machine_id,factory_id) VALUES(?,?,?)", (tid, mid, fac))
            c.commit()
    return {"ok": True}


@router.delete("/groups/{tid}/members/{mid}")
async def del_member(tid: int, mid: int, req: Request):
    require_role(req, "planner", "admin")
    with closing(db()) as c:
        _ensure(c)
        c.execute("DELETE FROM pm_members WHERE template_id=? AND machine_id=?", (tid, mid))
        c.commit()
    return {"ok": True}


@router.post("/assign")
async def assign(req: Request):
    """Pin a machine to a specific checklist sheet (overrides name matching, persists)."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    mid, tid = b.get("machine_id"), b.get("template_id")
    if not mid:
        raise HTTPException(400, "machine_id required")
    with closing(db()) as c:
        _ensure(c)
        c.execute("DELETE FROM pm_members WHERE machine_id=?", (mid,))
        if tid:
            c.execute("INSERT INTO pm_members(template_id,machine_id,factory_id) VALUES(?,?,?)", (tid, mid, fac))
        c.commit()
    return {"ok": True}


@router.post("/mark")
async def mark(req: Request):
    """Toggle one square (a Week/1M/3M/6M/1Y cell) for a machine's checklist item."""
    require_role(req, "planner", "admin", "technician")
    b = await req.json()
    mid, iid, cell, on = b.get("machine_id"), b.get("item_id"), b.get("cell"), bool(b.get("on"))
    if not (mid and iid and cell):
        raise HTTPException(400, "machine_id, item_id, cell required")
    with closing(db()) as c:
        _ensure(c)
        exists = c.execute("SELECT 1 FROM pm_marks WHERE machine_id=? AND item_id=? AND cell=?",
                           (mid, iid, cell)).fetchone()
        if on and not exists:
            c.execute("INSERT INTO pm_marks(machine_id,item_id,cell) VALUES(?,?,?)", (mid, iid, cell))
        elif not on and exists:
            c.execute("DELETE FROM pm_marks WHERE machine_id=? AND item_id=? AND cell=?", (mid, iid, cell))
        c.commit()
    return {"ok": True}


@router.patch("/items/{iid}")
async def edit_item(iid: int, req: Request):
    """Change a checklist item's frequency (the colour code on the sheet)."""
    u = require_role(req, "planner", "admin")
    b = await req.json()
    fq = b.get("freq")
    if fq not in FREQ_DAYS:
        raise HTTPException(400, "bad freq")
    with closing(db()) as c:
        _ensure(c)
        it = c.execute("SELECT i.*, t.factory_id fac FROM pm_items i JOIN pm_templates t ON t.id=i.template_id"
                       " WHERE i.id=?", (iid,)).fetchone()
        if not it:
            raise HTTPException(404, "item not found")
        if it["fac"] != _fac(u):
            raise HTTPException(403, "this checklist belongs to another plant")
        if (it["freq"] or "") != fq:
            c.execute("UPDATE pm_items SET freq=?, app_edit=1 WHERE id=?", (fq, iid))
            _item_log(c, it["fac"], it["template_id"], iid, "edit", _item_view(dict(it)),
                      _item_view({**dict(it), "freq": fq}), u)
        c.commit()
    return {"ok": True}


# ── editing a checklist in the app ──────────────────────────────────────────────
ITEM_FIELDS = ("item", "normal_status", "method", "freq")


def _item_view(r):
    return json.dumps({"seq": r.get("seq"), "item": r.get("item") or "",
                       "normal_status": r.get("normal_status") or "",
                       "method": r.get("method") or "", "freq": r.get("freq") or ""},
                      ensure_ascii=False)


def _item_log(c, fac, tid, iid, action, before, after, u):
    _ensure_records(c)
    c.execute("INSERT INTO pm_item_log(factory_id,template_id,item_id,action,before,after,user_id,user_name,at)"
              " VALUES(?,?,?,?,?,?,?,?,?)",
              (fac, tid, iid, action, before or "", after or "",
               (u or {}).get("id"), (u or {}).get("name") or (u or {}).get("username") or "", now()))


def _template_for_edit(c, tid, u):
    t = c.execute("SELECT * FROM pm_templates WHERE id=?", (tid,)).fetchone()
    if not t:
        raise HTTPException(404, "checklist not found")
    if t["factory_id"] != _fac(u):
        raise HTTPException(403, "ไม่ใช่เช็คลิสต์ของโรงงานนี้ / this checklist belongs to another plant")
    return t


def _template_users(c, fac, tid):
    """The machines this sheet is used by — the edit screen says how far a change reaches."""
    mids = [m for m, t in _match_all(c, fac).items() if t == tid]
    codes = []
    if mids:
        ph = ",".join("?" * len(mids[:2000]))
        codes = [r["code"] for r in c.execute(
            f"SELECT code FROM machines WHERE id IN ({ph}) ORDER BY code", mids[:2000])]
    return codes


@router.get("/templates/{tid}/usage")
async def template_usage(tid: int, req: Request):
    u = require_role(req, "planner", "admin", "manager")
    with closing(db()) as c:
        _ensure(c)
        t = _template_for_edit(c, tid, u)
        codes = _template_users(c, t["factory_id"], tid)
        started = c.execute(
            "SELECT COUNT(*) n FROM jobs j WHERE j.jobtype='PM' AND j.status IN ('InProgress','Paused','Hold','Rework')"
            " AND j.id IN (SELECT job_id FROM pm_job_items)"
            " AND j.machine_id IN (SELECT id FROM machines WHERE factory_id=?)", (t["factory_id"],)).fetchone()["n"]
    return {"template_id": tid, "machine_type": t["machine_type"], "machines": len(codes),
            "codes": codes[:12], "started_jobs_plant": started,
            "can_edit": u["role"] in ("planner", "admin")}


@router.get("/templates/{tid}/history")
async def template_history(tid: int, req: Request):
    u = require_role(req, "planner", "admin", "manager")
    with closing(db()) as c:
        _ensure(c)
        t = _template_for_edit(c, tid, u)
        rows = [dict(r) for r in c.execute(
            "SELECT * FROM pm_item_log WHERE (template_id=? OR (template_id IS NULL AND factory_id=?"
            " AND action='upload')) ORDER BY id DESC LIMIT 120", (tid, t["factory_id"]))]
    for r in rows:
        for k in ("before", "after"):
            try:
                r[k] = json.loads(r[k]) if r[k] else None
            except Exception:
                pass
    return {"template_id": tid, "machine_type": t["machine_type"], "rows": rows}


@router.post("/templates/{tid}/items")
async def save_template_items(tid: int, req: Request):
    """Save an edited checklist: the whole list, in order, as the edit screen holds it.

    Body {items:[{id?, item, normal_status, method, freq, removed?}]}. An item with an id
    is updated (or hidden when removed); one without is added. Items are never deleted —
    a removed one is hidden from future PM and still answers for the records that used
    it. Jobs already started keep their own frozen list; jobs not started pick this up.
    Refused if the sheet changed under the editor (another person saved first).
    """
    u = require_role(req, "planner", "admin")
    b = await req.json()
    posted = b.get("items")
    if not isinstance(posted, list):
        raise HTTPException(400, "items must be a list")
    with closing(db()) as c:
        _ensure(c)
        t = _template_for_edit(c, tid, u)
        fac = t["factory_id"]
        cur = {r["id"]: dict(r) for r in c.execute(
            "SELECT * FROM pm_items WHERE template_id=?" + ACTIVE_ITEM, (tid,))}
        ids = [int(x["id"]) for x in posted if str(x.get("id") or "").isdigit()]
        if set(cur) - set(ids) or any(i not in cur for i in ids):
            raise HTTPException(409, "เช็คลิสต์นี้ถูกแก้ไขโดยคนอื่นแล้ว — โหลดใหม่ก่อน /"
                                     " someone else changed this checklist — reload it first")
        seq = 0
        n_edit = n_add = n_rm = 0
        for x in posted:
            iid = int(x["id"]) if str(x.get("id") or "").isdigit() else None
            txt = str(x.get("item") or "").strip()[:400]
            nrm = str(x.get("normal_status") or "").strip()[:200]
            mth = str(x.get("method") or "").strip()[:40]
            fq = x.get("freq") or "monthly"
            if fq not in FREQ_DAYS:
                raise HTTPException(400, f"bad frequency: {fq}")
            if iid and x.get("removed"):
                c.execute("UPDATE pm_items SET active=0, app_edit=1 WHERE id=?", (iid,))
                _item_log(c, fac, tid, iid, "remove", _item_view(cur[iid]), "", u)
                n_rm += 1
                continue
            if not txt:
                if iid:
                    raise HTTPException(400, "an item cannot be left empty — remove it instead")
                continue
            seq += 1
            if iid:
                old = cur[iid]
                new = {"seq": seq, "item": txt, "normal_status": nrm, "method": mth, "freq": fq}
                changed = any((old.get(k) or "") != new[k] for k in ITEM_FIELDS)
                if changed or old.get("seq") != seq:
                    c.execute("UPDATE pm_items SET seq=?,item=?,normal_status=?,method=?,freq=?"
                              + (",app_edit=1" if changed else "") + " WHERE id=?",
                              (seq, txt, nrm, mth, fq, iid))
                if changed:
                    _item_log(c, fac, tid, iid, "edit", _item_view(old), _item_view(new), u)
                    n_edit += 1
            else:
                nid = c.insert_id("INSERT INTO pm_items(template_id,seq,item,normal_status,method,freq,img,active,app_edit)"
                                  " VALUES(?,?,?,?,?,?,'',1,1)", (tid, seq, txt, nrm, mth, fq))
                _item_log(c, fac, tid, nid, "add", "",
                          _item_view({"seq": seq, "item": txt, "normal_status": nrm, "method": mth, "freq": fq}), u)
                n_add += 1
        if not seq:
            raise HTTPException(400, "a checklist needs at least one item")
        c.commit()
        _MATCH_CACHE.pop(fac, None)
        items = [dict(r) for r in c.execute(
            "SELECT id,seq,item,normal_status,method,freq,img FROM pm_items WHERE template_id=?"
            + ACTIVE_ITEM + " ORDER BY seq", (tid,))]
    return {"ok": True, "edited": n_edit, "added": n_add, "removed": n_rm, "items": items}


@router.get("/machine-template")
async def machine_template(req: Request):
    """The PM sheet for one asset: the template whose type best matches the machine
    (overridable), rendered from the imported workbook data (items + freq colour + img)."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    mid = req.query_params.get("machine_id")
    force = req.query_params.get("template_id")
    with closing(db()) as c:
        _ensure(c)
        m = c.execute("SELECT id,code,name,asset_group,category FROM machines WHERE id=? AND factory_id=?",
                      (mid, fac)).fetchone()
        if not m:
            raise HTTPException(404, "machine not found")
        hay = " ".join([m["name"] or "", m["asset_group"] or "", m["category"] or ""]).lower()
        an = _norm(m["name"])
        scored = []
        for t in c.execute("SELECT id,machine_type FROM pm_templates WHERE factory_id=?", (fac,)):
            toks = [x.lower() for x in _toks(t["machine_type"])]
            sc = sum(1 for x in toks if x in hay)
            tn = _norm(t["machine_type"])
            if tn and (tn in an or an in tn):        # exact normalized name match wins
                sc += 1000 + len(tn)
            sc += _alias_bonus(t["machine_type"], an)
            scored.append((sc, t["id"], t["machine_type"]))
        scored.sort(key=lambda z: -z[0])
        member_ids = [r["template_id"] for r in c.execute(
            "SELECT template_id FROM pm_members WHERE machine_id=?", (mid,))]
        members = [{"id": t["id"], "machine_type": t["machine_type"], "score": 99, "member": True}
                   for t in c.execute("SELECT id,machine_type FROM pm_templates WHERE factory_id=?", (fac,))
                   if t["id"] in member_ids]
        namematch = [{"id": z[1], "machine_type": z[2], "score": z[0]} for z in scored[:10]
                     if z[0] > 0 and z[1] not in member_ids]
        candidates = members + namematch
        chosen = None
        if force:
            chosen = c.execute("SELECT id,machine_type,form_no FROM pm_templates WHERE id=? AND factory_id=?",
                               (force, fac)).fetchone()
        if not chosen and members:
            chosen = c.execute("SELECT id,machine_type,form_no FROM pm_templates WHERE id=?", (members[0]["id"],)).fetchone()
        if not chosen and scored and scored[0][0] > 0:
            chosen = c.execute("SELECT id,machine_type,form_no FROM pm_templates WHERE id=?", (scored[0][1],)).fetchone()
        tpl = None
        if chosen:
            items = [dict(x) for x in c.execute(
                "SELECT id,seq,item,normal_status,method,freq,img FROM pm_items WHERE template_id=?"
                + ACTIVE_ITEM + " ORDER BY seq", (chosen["id"],))]
            marks = {}
            for r in c.execute("SELECT item_id,cell FROM pm_marks WHERE machine_id=?", (mid,)):
                marks.setdefault(r["item_id"], []).append(r["cell"])
            for it in items:
                it["marks"] = marks.get(it["id"], [])
            tpl = {"id": chosen["id"], "machine_type": chosen["machine_type"],
                   "form_no": chosen["form_no"] or "F-SP-ENG02-01 Rev.01", "items": items}
    return {"machine": {"id": m["id"], "code": m["code"], "name": m["name"]},
            "template": tpl, "matched": bool(chosen), "candidates": candidates}


CELL_FREQ = {"w1": "weekly", "w2": "weekly", "w3": "weekly", "w4": "weekly", "w5": "weekly",
             "1M": "monthly", "3M": "q3m", "6M": "m6", "1Y": "yearly"}
FREQ_STEP = {"weekly": ("d", 7), "monthly": ("m", 1), "q3m": ("m", 3), "m6": ("m", 6), "yearly": ("m", 12)}
FREQ_ORD = ["weekly", "monthly", "q3m", "m6", "yearly"]


def _add_months(d, n):
    from calendar import monthrange
    mo = d.month - 1 + n
    y = d.year + mo // 12
    mo = mo % 12 + 1
    return date(y, mo, min(d.day, monthrange(y, mo)[1]))


def _gen(start, frm, to, freq):
    kind, step = FREQ_STEP.get(freq, (None, 0))
    if not kind or to < start:
        return []
    res, k = [], 0
    while k < 4000:
        d = (start + timedelta(days=step * k)) if kind == "d" else _add_months(start, step * k)
        if d > to:
            break
        if d >= frm and d >= start:
            res.append(d.isoformat())
        k += 1
    return res


def _gen_weekly(start, frm, to, dow):
    """Weekly occurrences on a specific weekday (dow 0=Mon..5=Sat), so weekly PM is
    spread across the 6 working days instead of all on the start weekday."""
    monday = start - timedelta(days=start.weekday())
    d = monday + timedelta(days=dow)
    if d < start:
        d += timedelta(days=7)
    res, k = [], 0
    while d <= to and k < 400:
        if d >= frm:
            res.append(d.isoformat())
        d += timedelta(days=7)
        k += 1
    return res


def _shift_sunday(iso):
    d = date.fromisoformat(iso)
    return (d + timedelta(days=1)).isoformat() if d.weekday() == 6 else iso


def _add_workdays(d, n):
    """Advance n working days (skip Sundays) — used to stagger monthly/quarterly PM
    across the days of the period instead of all on the start day."""
    while n > 0:
        d += timedelta(days=1)
        if d.weekday() != 6:
            n -= 1
    return d


# ── the route ──────────────────────────────────────────────────────────────────
# A tower is expensive to enter. The crew carries its tools up, and coming back
# down only to climb again is the most wasted hour in the week — so the towers are
# walked first, together, and the descent to ground level happens once. Everything
# that is not a tower is ground work and sorts after it.
CLIMB_RE = re.compile(r"tower|หอคอย", re.I)   # "tower" / หอคอย (tower)

# How many weeks one cycle of each frequency spans. Used as a WEIGHT: a monthly
# machine asks for a quarter of the attention a weekly one does, a yearly machine a
# fiftieth. Weighting is what lets every frequency share one route instead of each
# walking its own.
FREQ_WEEKS = {"weekly": 1.0, "monthly": 4.33, "q3m": 13.0, "m6": 26.0, "yearly": 52.0}
FREQ_MONTHS = {"monthly": 1, "q3m": 3, "m6": 6, "yearly": 12}
WEEK_SLOTS = 4                      # every month holds at least four of each weekday


def _area(line):
    """The building a machine stands in — the first segment of its location."""
    return (line or "").split("/")[0].strip()


# The floor is the second climb. A tower day that hops 1 -> 6 -> 2 -> 8 costs the crew
# the same stairs twice, so the walk is ordered by floor inside each area. BFL keeps its
# floor in machines.floor; BFLPC keeps it inside the location string as "ชั้น N"
# (chan N = floor N), so both are read and either will do.
FLOOR_RE = re.compile(r"ชั้น\s*([^\s/]+)|\bfloor\s*([A-Za-z]?\d+)", re.I)


def _floor_of(floor, line):
    """The floor a machine stands on, as it is written — '2', 'B1', or '' if unknown."""
    f = (floor or "").strip()
    if f:
        return f
    m = FLOOR_RE.search(line or "")
    return ((m.group(1) or m.group(2) or "").strip() if m else "")


def _floor_sort(f):
    """Ground up: 1, 2, ... 12, then the lettered levels (A1, B1), then unknown last."""
    f = (f or "").strip()
    if not f:
        return (2, "", 0)
    if f.isdigit():
        return (0, "", int(f))
    head = f[0].upper()
    tail = "".join(ch for ch in f[1:] if ch.isdigit())
    return (1, head, int(tail) if tail else 0)


def _route_key(line, code, floor=""):
    """Order the whole plant into one walk: towers first, then ground; inside each the
    areas in order, and inside an area the floors from the bottom up — so a room is
    finished before the next is started and a floor before the next is climbed. A
    machine with no location sorts last, so unfilled rows collect in one place instead
    of salting every day."""
    a = _area(line)
    return (0 if (a and CLIMB_RE.search(a)) else 1, (a or "~~~").lower(),
            _floor_sort(floor), (line or "~~~").lower(), code or "")


def _nth_weekday(y, m, dow, n):
    """The (n+1)-th <dow> of that month, or None if the month has no such day.
    Monthly PM anchors to this rather than to a date, because a fixed date lands on
    a different weekday every month and would send the crew up the wrong tower."""
    d = date(y, m, 1)
    d += timedelta(days=(dow - d.weekday()) % 7 + 7 * n)
    return d if d.month == m else None


def _cal_config(c, fac):
    """(weekly_off set of weekday ints, holidays set of ISO dates) for this factory."""
    r = c.execute("SELECT hours_json FROM factories WHERE id=?", (fac,)).fetchone()
    try:
        cfg = json.loads(r["hours_json"]) if r and r["hours_json"] else {}
    except Exception:
        cfg = {}
    # offdays_of, not weekly_off: a weekday given its own hours (Sunday 07:00-20:00)
    # is a WORKING day, and PM must be plannable on it. The two would disagree
    # otherwise — the KPI clock counting Sunday hours the PM calendar refuses to use.
    off = set(offdays_of(cfg))
    return off, set(cfg.get("holidays", []) or [])


def _next_workday(d, offdays, holidays):
    """Nudge a date off weekly-off days and holidays to the next working day."""
    for _ in range(400):
        if d.weekday() not in offdays and d.isoformat() not in holidays:
            return d
        d += timedelta(days=1)
    return d


def _pm_config(c, fac):
    r = c.execute("SELECT start_date,active FROM pm_config WHERE factory_id=?", (fac,)).fetchone()
    return ((r["start_date"] if r else "") or "", (r["active"] if r else 0))


# Matching every machine to its PM sheet is a fuzzy name comparison across the whole
# register — 151 machines against 84 templates, ~600ms a call on the FP plant. That was
# fine while only the calendar asked for it once a page. The technician's checklist asks
# on every job open and again on every Stop, and 600ms twice over a factory wifi is felt
# in the hand. So the map is cached against a cheap fingerprint of the things that could
# change it: how many templates, items, pins and machines there are, and the total
# length of the machine names, which catches a rename that leaves the counts alone.
# Any edit moves the fingerprint and the next call rebuilds; nothing goes stale.
_MATCH_CACHE = {}


def _match_sig(c, fac):
    def one(q, a=()):
        r = c.execute(q, a).fetchone()
        return (r[0] if r and r[0] is not None else 0)
    return (one("SELECT COUNT(*) FROM pm_members WHERE factory_id=?", (fac,)),
            one("SELECT COUNT(*) FROM pm_templates WHERE factory_id=?", (fac,)),
            one("SELECT COUNT(*) FROM pm_items"),
            one("SELECT COALESCE(SUM(COALESCE(active,1)),0) FROM pm_items"),
            one("SELECT COUNT(*) FROM machines WHERE factory_id=?", (fac,)),
            one("SELECT COALESCE(SUM(LENGTH(name)),0) FROM machines WHERE factory_id=?", (fac,)),
            one("SELECT COALESCE(SUM(LENGTH(machine_type)),0) FROM pm_templates WHERE factory_id=?", (fac,)))


def _match_all(c, fac):
    """Cached wrapper — see _MATCH_CACHE above. Same answer, ~0ms when nothing moved."""
    try:
        sig = _match_sig(c, fac)
    except Exception:
        return _match_all_uncached(c, fac)
    hit = _MATCH_CACHE.get(fac)
    if hit and hit[0] == sig:
        return hit[1]
    out = _match_all_uncached(c, fac)
    _MATCH_CACHE[fac] = (sig, out)
    return out


def _match_all_uncached(c, fac):
    """{machine_id: template_id} — explicit assignment first, else normalized name /
    alias match against the sheet names."""
    tpls = [(t["id"], [x.lower() for x in _toks(t["machine_type"])], _norm(t["machine_type"]), t["machine_type"])
            for t in c.execute("SELECT id,machine_type FROM pm_templates WHERE factory_id=?", (fac,))]
    members = {r["machine_id"]: r["template_id"]
               for r in c.execute("SELECT machine_id,template_id FROM pm_members WHERE factory_id=?", (fac,))}
    out = {}
    for m in c.execute("SELECT id,code,name,asset_group,category FROM machines WHERE factory_id=?", (fac,)):
        tid = members.get(m["id"])
        if not tid:
            hay = " ".join([m["name"] or "", m["asset_group"] or "", m["category"] or "", m["code"] or ""]).lower()
            an = _norm(m["name"])
            scored = []
            for t_id, toks, tn, tmt in tpls:
                s = sum(1 for k in toks if k in hay)
                if tn and tn == an:
                    # the machine IS this sheet — must beat a longer sheet name that
                    # merely contains it (เครื่องตัดปลา vs เครื่องตัดปลาชาดีน)
                    s += 5000 + len(tn)
                elif tn and (tn in an or an in tn):
                    s += 1000 + len(tn)
                s += _alias_bonus(tmt, an)
                if s > 0:
                    scored.append((s, t_id, tn, tmt))
            tid = None
            if scored:
                top_s, top_id, top_tn, top_tmt = max(scored, key=lambda x: x[0])
                # Sheets split per machine — เครื่องบด / เครื่องบด 3, Hi-Mixer / HIMIXER3,
                # TTK-1 / TTK-2 — collapse to the same normalised name, so they are one
                # family. Inside a family the machine's own number picks its copy; a
                # machine with no copy of its own falls back to the plain sheet. This is
                # decided on the family, not on the score, because a stray shared word
                # can otherwise hand the machine to the wrong copy by a single point.
                # A copy may also be named shorter than the sheet it splits off from
                # (Ribbon Mixer / Ribbon3), so a numbered sheet whose name is the start
                # of the winner's name counts as family too.
                fam = [(t_id, tmt) for s, t_id, tn, tmt in scored
                       if top_tn and tn and (tn == top_tn
                                             or (_tail_no(tmt) is not None and top_tn.startswith(tn)))]
                if len(fam) > 1:
                    anum = _tail_no(m["name"])
                    same = [t for t in fam if anum is not None and _tail_no(t[1]) == anum]
                    if same:
                        fam = same
                    else:
                        plain = [t for t in fam if _tail_no(t[1]) is None]
                        if plain:
                            fam = plain
                    tid = fam[0][0]
                else:
                    tid = top_id
        if tid:
            out[m["id"]] = tid
    return out


def _machine_freqs(c, fac):
    tfreq = {}
    for r in c.execute("""SELECT t.id tid, i.freq FROM pm_templates t JOIN pm_items i ON i.template_id=t.id
            WHERE t.factory_id=? AND COALESCE(i.active,1)=1 GROUP BY t.id, i.freq""", (fac,)):
        tfreq.setdefault(r["tid"], set()).add(r["freq"])
    m2t = _match_all(c, fac)
    minfo = {m["id"]: m for m in
             c.execute("SELECT id,code,name,line,floor FROM machines WHERE factory_id=?", (fac,))}
    out = {}
    for mid, tid in m2t.items():
        freqs = {f for f in tfreq.get(tid, set()) if f in FREQ_STEP}
        if freqs and mid in minfo:
            m = minfo[mid]
            out[mid] = {"code": m["code"], "name": m["name"], "line": m["line"] or "",
                        "floor": _floor_of(m["floor"], m["line"]), "freqs": freqs}
    return out


# Colours for a group a planner invents on the Assets page. The PM sheets bring their
# own tab colours; these are the ones left over, so a new area group is still a colour
# the crew can be told to walk — which is how the groups get used.
NEW_GROUP_COLORS = ["FF0EA5E9", "FF14B8A6", "FFF97316", "FFA855F7", "FFEC4899",
                    "FF84CC16", "FF6366F1", "FFEAB308", "FF64748B", "FF0891B2"]


def _free_group_color(c, fac):
    """The first palette colour no template and no group is using yet."""
    used = {(r["color"] or "").upper() for r in
            c.execute("SELECT color FROM pm_templates WHERE factory_id=?", (fac,))}
    used |= {(r["color"] or "").upper() for r in
             c.execute("SELECT color FROM pm_groups WHERE factory_id=?", (fac,))}
    for col in NEW_GROUP_COLORS:
        if col not in used:
            return col
    n = 1
    while f"FFUSER{n:02d}"[:8] in used:
        n += 1
    return f"FFUSER{n:02d}"[:8]


@router.get("/color-groups")
async def color_groups(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        tcolor = {t["id"]: (t["color"] or "") for t in c.execute("SELECT id,color FROM pm_templates WHERE factory_id=?", (fac,))}
        gname = {r["color"]: r["name"] for r in c.execute("SELECT color,name FROM pm_groups WHERE factory_id=?", (fac,))}
        m2t = _match_all(c, fac)
        # a colour pinned on the asset wins over the one its PM sheet brings
        ovr = {m["id"]: (m["pm_group_color"] or "")
               for m in c.execute("SELECT id,pm_group_color FROM machines WHERE factory_id=?", (fac,))}
    from collections import Counter
    tpl_per = Counter(v for v in tcolor.values() if v)
    ast_per = Counter()
    for mid, pinned in ovr.items():
        col = pinned or tcolor.get(m2t.get(mid), "")
        if col:
            ast_per[col] += 1
    # groups with no template of their own (made by hand on the Assets page) belong here too
    colors = sorted(set(tpl_per) | set(gname), key=lambda x: (-tpl_per.get(x, 0), -ast_per.get(x, 0), x))
    return {"groups": [{"color": c, "name": gname.get(c, c), "n_templates": tpl_per.get(c, 0),
                        "n_assets": ast_per.get(c, 0)} for c in colors]}


@router.patch("/color-groups")
async def rename_group(req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    color, name = b.get("color"), (b.get("name") or "").strip()
    if not color:
        raise HTTPException(400, "color required")
    with closing(db()) as c:
        _ensure(c)
        if c.execute("SELECT 1 FROM pm_groups WHERE factory_id=? AND color=?", (fac, color)).fetchone():
            c.execute("UPDATE pm_groups SET name=? WHERE factory_id=? AND color=?", (name, fac, color))
        else:
            c.execute("INSERT INTO pm_groups(factory_id,color,name) VALUES(?,?,?)", (fac, color, name))
        c.commit()
    return {"ok": True}


ASSET_COLS = ("id,code,name,category,asset_group,criticality,line,floor,department,"
              "manufacturer,brand_model,serial_no,size,year_install,last_pm_date,"
              "pm_freq_days,active,remark,pm_group_color")


@router.get("/assets")
async def list_assets(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        m2t = _match_all(c, fac)
        tinfo = {t["id"]: t for t in c.execute("SELECT id,machine_type,color FROM pm_templates WHERE factory_id=?", (fac,))}
        gname = {r["color"]: r["name"] for r in c.execute("SELECT color,name FROM pm_groups WHERE factory_id=?", (fac,))}
        rows = []
        # every column the asset register holds — the page shows a short set by default
        # and the full one on demand, and the Excel export writes all of it
        for m in c.execute(f"SELECT {ASSET_COLS} FROM machines WHERE factory_id=? ORDER BY code",
                           (fac,)):
            tid = m2t.get(m["id"])
            t = tinfo.get(tid) if tid else None
            # a colour set by hand on the row beats the one the template brings
            color = (m["pm_group_color"] or "") or ((t["color"] if t else "") or "")
            rows.append({**dict(m), "pm_sheet": (t["machine_type"] if t else ""),
                         "pm_group": gname.get(color, color) if color else "", "pm_color": color})
        # every sheet name in the factory, so the edit row can offer them all —
        # not only the ones some asset already matches
        sheets = sorted({(t["machine_type"] or "") for t in tinfo.values() if t["machine_type"]})
    return {"assets": rows, "sheets": sheets}


@router.get("/assets/export")
async def export_assets(req: Request, q: str = "", tab: str = "all"):
    """The asset register as a real .xlsx, honouring the tab and search on screen."""
    import io
    from fastapi.responses import StreamingResponse
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl is not installed on the server")

    data = await list_assets(req)
    rows = data["assets"]
    if tab == "crit":
        rows = [a for a in rows if (a.get("criticality") or "").upper() == "A"]
    elif tab == "unmatched":
        rows = [a for a in rows if not a.get("pm_group")]
    if q:
        ql = q.lower()
        rows = [a for a in rows
                if ql in (a.get("code") or "").lower() or ql in (a.get("name") or "").lower()]

    heads = [("Asset ID", "code", 16), ("Name", "name", 40), ("Category", "category", 16),
             ("Group", "asset_group", 22), ("Criticality", "criticality", 11),
             ("PM group", "pm_group", 18), ("PM sheet", "pm_sheet", 26),
             ("Line / location", "line", 26), ("Floor", "floor", 10),
             ("Department", "department", 18), ("Manufacturer", "manufacturer", 20),
             ("Brand / model", "brand_model", 22), ("Serial no", "serial_no", 18),
             ("Size", "size", 14), ("Year installed", "year_install", 13),
             ("Last PM", "last_pm_date", 13), ("PM every (days)", "pm_freq_days", 14),
             ("Active", "active", 8), ("Remark", "remark", 30)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    hfill = PatternFill("solid", fgColor="1E293B")
    hfont = Font(bold=True, color="FFFFFF", size=10)
    for i, (label, _k, w) in enumerate(heads, start=1):
        cell = ws.cell(row=1, column=i, value=label)
        cell.fill, cell.font = hfill, hfont
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, a in enumerate(rows, start=2):
        for i, (_label, k, _w) in enumerate(heads, start=1):
            v = a.get(k, "")
            if k == "active":
                v = "yes" if v else "no"
            ws.cell(row=r, column=i, value=v)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(1, len(rows) + 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"assets-{today()}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.post("/assets")
async def add_asset(req: Request):
    """Create a new asset in this factory (same fields the Excel import uses)."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    code = (b.get("code") or "").strip()
    if not code:
        raise HTTPException(400, "asset ID (code) required")
    with closing(db()) as c:
        _ensure(c)
        if c.execute("SELECT 1 FROM machines WHERE code=? AND factory_id=?", (code, fac)).fetchone():
            raise HTTPException(400, "asset ID already exists")
        c.execute("""INSERT INTO machines(code,name,category,asset_group,criticality,factory_id,active)
            VALUES(?,?,?,?,?,?,1)""", (code, b.get("name", ""), b.get("category", ""),
                                       b.get("asset_group", ""), b.get("criticality", ""), fac))
        c.commit()
    return {"ok": True}


# The plain columns the Assets page writes straight onto the machine row. PM sheet and
# PM group are handled separately below: the sheet pins a pm_members row (the explicit
# template assignment _match_all already prefers), the group writes a colour override.
ASSET_EDIT = ("code", "name", "category", "asset_group", "criticality", "line", "floor",
              "department", "manufacturer", "brand_model", "serial_no", "size",
              "year_install", "last_pm_date", "remark")


@router.patch("/assets/{mid}")
async def edit_asset(mid: int, req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    sets = {k: ("" if b[k] is None else str(b[k]).strip()) for k in ASSET_EDIT if k in b}
    sheet = None if b.get("pm_sheet") is None else str(b["pm_sheet"]).strip()
    group = None if b.get("pm_group") is None else str(b["pm_group"]).strip()
    if not sets and sheet is None and group is None:
        raise HTTPException(400, "nothing to update")
    with closing(db()) as c:
        _ensure(c)
        row = c.execute("SELECT code FROM machines WHERE id=? AND factory_id=?",
                        (mid, fac)).fetchone()
        if not row:
            raise HTTPException(404, "ไม่พบเครื่องนี้ในโรงงานนี้ / asset not found in this factory")
        if "code" in sets:
            # the Asset ID is how every job, PM plan and report names the machine,
            # so it has to stay filled in and unique inside the factory
            if not sets["code"]:
                raise HTTPException(400, "รหัสเครื่องต้องไม่ว่าง / asset ID cannot be blank")
            if sets["code"] != row["code"] and c.execute(
                    "SELECT 1 FROM machines WHERE code=? AND factory_id=? AND id<>?",
                    (sets["code"], fac, mid)).fetchone():
                raise HTTPException(400, f"รหัส {sets['code']} มีอยู่แล้ว / asset ID already in use")
        # both PM columns are resolved before anything is written, so a typo cannot
        # leave the asset half-moved
        tid = None
        if sheet:
            t = c.execute("SELECT id FROM pm_templates WHERE factory_id=? AND lower(machine_type)=lower(?)",
                          (fac, sheet)).fetchone()
            if not t:
                raise HTTPException(400, f"ไม่พบชีต PM ชื่อ {sheet} / no PM sheet named {sheet}")
            tid = t["id"]
        color = ""
        new_group = ""
        if group:
            g = c.execute("SELECT color FROM pm_groups WHERE factory_id=? AND lower(name)=lower(?)",
                          (fac, group)).fetchone()
            if not g:
                g = c.execute("SELECT color FROM pm_templates WHERE factory_id=? AND lower(color)=lower(?) LIMIT 1",
                              (fac, group)).fetchone()
            if g:
                color = g["color"] or ""
            else:
                # a name nobody has used yet starts a new group on the next free colour;
                # rename it later from the Colour groups row, e.g. to an area name
                color = _free_group_color(c, fac)
                c.execute("INSERT INTO pm_groups(factory_id,color,name) VALUES(?,?,?)",
                          (fac, color, group))
                new_group = group
        if sets:
            c.execute(f"UPDATE machines SET {','.join(k+'=?' for k in sets)} WHERE id=? AND factory_id=?",
                      (*sets.values(), mid, fac))
        if sheet is not None:
            # blank clears the pin and the asset goes back to name matching
            c.execute("DELETE FROM pm_members WHERE machine_id=? AND factory_id=?", (mid, fac))
            if tid:
                c.execute("INSERT INTO pm_members(template_id,machine_id,factory_id) VALUES(?,?,?)",
                          (tid, mid, fac))
        if group is not None:
            c.execute("UPDATE machines SET pm_group_color=? WHERE id=? AND factory_id=?",
                      (color, mid, fac))
        c.commit()
    return {"ok": True, "new_group": new_group}


STARTED = ("InProgress", "Paused")


def day_counts(c, fac, d_from, d_to, occ=None):
    """One day's work, counted the SAME way on every screen that shows it.

    The calendar and the assign board each counted a day by their own rule, and the two
    never agreed: the calendar counted PM from the programme (finished ones included) but
    finished CM not at all, and folded duplicate work orders into one; the board counted
    only what is still open, plus work carried in from other days and work with no date.
    A planner looking at 76 on one and 43 on the other had no way to tell which was
    right. Both now read this.

    Every work order PLANNED for the day, in exactly one bucket, plus the PM the
    programme puts on that day that has no work order yet (the board's "DUE" cards, by
    the same rule as `teams._pm_due`):

        done      finished and accepted                (Done)
        accept    finished, waiting for acceptance     (ServiceCompleted)
        working   a technician has started it          (InProgress, Paused)
        assigned  has a technician, not started yet
        waiting   no technician yet, including PM due with no work order
        total     the sum of the five

    Cancelled and rejected work is not work. Work carried in from an earlier day and
    work with no date belong to no day, so they are not in here; the board shows them
    separately beside this total.
    """
    out = {}
    zero = {"total": 0, "done": 0, "accept": 0, "working": 0, "assigned": 0, "waiting": 0}

    def slot(d):
        return out.setdefault(d, dict(zero))
    for r in c.execute(
            """SELECT j.planned_date d, j.status s, j.lead_tech lt
                 FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id
                WHERE j.planned_date BETWEEN ? AND ? AND j.status NOT IN ('Cancelled','Rejected')
                  AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)""",
            (d_from, d_to, fac)):
        s = r["s"]
        key = ("done" if s == "Done" else "accept" if s == "ServiceCompleted"
               else "working" if s in STARTED else "assigned" if r["lt"] else "waiting")
        k = slot(str(r["d"])[:10])
        k[key] += 1
        k["total"] += 1
    if occ is None:
        try:
            occ = _occurrences(c, fac, date.fromisoformat(d_from), date.fromisoformat(d_to))
        except Exception:
            occ = []
    have = set()
    for r in c.execute(
            "SELECT machine_id, planned_date, due_date FROM jobs WHERE jobtype='PM'"
            " AND machine_id IS NOT NULL AND status NOT IN ('Cancelled','Rejected')"
            " AND (planned_date BETWEEN ? AND ? OR due_date BETWEEN ? AND ?)",
            (d_from, d_to, d_from, d_to)):
        for v in (r["planned_date"], r["due_date"]):
            if v:
                have.add((r["machine_id"], str(v)[:10]))
    for o in occ:
        if o.get("cm"):
            continue
        d = str(o.get("date") or "")[:10]
        if not d or (o["machine_id"], d) in have:
            continue
        k = slot(d)
        k["waiting"] += 1
        k["total"] += 1
    return out


def _occurrences(c, fac, frm, to):
    start_s, active = _pm_config(c, fac)
    if not active or not start_s:
        return []
    try:
        start = date.fromisoformat(start_s[:10])
    except Exception:
        return []
    permac = _machine_freqs(c, fac)
    offdays, holidays = _cal_config(c, fac)
    # ── one route for the whole plant, shared by every frequency ────────────────
    #
    # Each frequency used to plan its own walk: weekly dealt the machines across six
    # days, monthly dealt the SAME machines across twenty-six, and the two never
    # agreed — so the weekly round was in Tower A while the monthly round was in the
    # Liquid Room, and the crew climbed a tower and came down again every day.
    # Measured on the BFLPC register that put them in both towers on 138 days of the
    # year. Worse, 3-month/6-month/yearly work was not spread at all: every machine
    # of that frequency fell on ONE day, giving a 459-job day against a normal 70.
    #
    # Now there is a single ordered walk — towers first, then ground, area by area —
    # and every piece of work, whatever its frequency, is dealt into the working days
    # of the week along that one walk. A machine is weighted by how often it comes
    # round (weekly 1, monthly a quarter, yearly a fiftieth), so each day carries an
    # equal share of the real load: the towers earn whole days, the ground areas
    # share what is left, and the descent happens once.
    wdays = [w for w in range(7) if w not in offdays] or [0, 1, 2, 3, 4, 5]
    nd = len(wdays)
    # Frequency last in the sort so that when one machine carries both a weekly and a
    # monthly checklist and they fall on the same day, the weekly row is the one that
    # survives the (day, machine) de-duplication below. One visit, and it is labelled
    # by the sheet the technician actually works from.
    work = sorted(((mid, fq) for mid, info in permac.items()
                   for fq in info["freqs"] if fq in FREQ_WEEKS),
                  key=lambda t: (_route_key(permac[t[0]]["line"], permac[t[0]]["code"],
                                            permac[t[0]].get("floor", "")),
                                 FREQ_ORD.index(t[1]) if t[1] in FREQ_ORD else 9))
    wof = lambda fq: 1.0 / FREQ_WEEKS[fq]
    load = sum(wof(fq) for _m, fq in work) or 1.0
    # Group the walk by AREA and give each area a span of the week measured in days.
    # Dealing machine by machine put the boundaries wherever the arithmetic fell:
    # Tower A's load came to 1.989 days, so one machine spilled onto the Tower B day
    # and sent the crew up the wrong tower for a single job. An area boundary that
    # lands within a quarter-day of a day edge is therefore snapped to it — the day
    # loads move by a job or two, and a tower is entered on the days it owns and no
    # others.
    order, groups = [], {}
    for mid, fq in work:
        a = _route_key(permac[mid]["line"], permac[mid]["code"])[1]     # the AREA, not the floor
        if a not in groups:
            order.append(a)
            groups[a] = []
        groups[a].append((mid, fq))
    bounds, pos = [0.0], 0.0
    for a in order:
        pos += sum(wof(fq) for _m, fq in groups[a]) / load * nd
        bounds.append(pos)
    bounds[-1] = float(nd)
    SNAP = 0.25
    for i in range(1, len(bounds) - 1):
        r = float(round(bounds[i]))
        if 0.0 < r < nd and abs(bounds[i] - r) <= SNAP and bounds[i - 1] <= r <= bounds[i + 1]:
            bounds[i] = r
    slot = {}
    for i, a in enumerate(order):
        lo, hi = bounds[i], bounds[i + 1]
        aload = sum(wof(fq) for _m, fq in groups[a]) or 1.0
        cum = 0.0
        for mid, fq in groups[a]:
            slot[(mid, fq)] = max(0, min(nd - 1, int(lo + (hi - lo) * cum / aload)))
            cum += wof(fq)
        # Every frequency of one machine shares a weekday, which is the point: the crew
        # is in that area on that day. When a machine's quarterly falls on its weekly
        # day the two are ONE visit, and the "longest sheet wins" rule below makes it
        # the quarterly — a quarterly service covers the weekly checks, so nothing is
        # lost and the machine is not climbed to twice. An earlier version pushed the
        # longer sheet onto another day of the same area to keep both rows; measured
        # over a year that changed no frequency's count materially (q3m 614 vs 623,
        # 6-month 644 vs 617) and cost the floor grouping its edge — it scattered one
        # or two jobs from floors 5, 7 and 8 across a day that was otherwise floors 1
        # and 2. Median floors per day went from 5 to 7 for nothing. It is gone.
    # Inside one day, the long-cycle work is spread across the weeks of its own cycle
    # — the 1st Tuesday, the 2nd Tuesday, and so on — and across the months of a
    # quarter or a year. The weekday never moves, so the crew's week keeps its shape
    # while a quarterly machine still comes round once a quarter.
    grp = {}
    for mid, fq in work:
        grp.setdefault((fq, slot[(mid, fq)]), []).append(mid)
    phase = {}
    for (fq, s), lst in grp.items():
        nmon = FREQ_MONTHS.get(fq, 0)
        for i, mid in enumerate(lst):
            phase[(mid, fq)] = (i % WEEK_SLOTS,
                                (i // WEEK_SLOTS) % nmon if nmon > 1 else 0)
    ov = {}
    for r in c.execute("""SELECT o.machine_id,o.orig_date,o.new_date FROM pm_overrides o
            JOIN machines m ON m.id=o.machine_id WHERE m.factory_id=?""", (fac,)):
        ov[(r["machine_id"], r["orig_date"])] = r["new_date"]
    jobstat = {}
    # a cancelled PM must not overwrite the real one on the same machine and day
    for j in c.execute("SELECT machine_id,planned_date,due_date,status FROM jobs"
                       " WHERE jobsource='PM' AND machine_id IS NOT NULL"
                       " AND status NOT IN " + VOID_SQL):
        d = (j["planned_date"] or j["due_date"] or "")[:10]
        if d:
            jobstat[(j["machine_id"], d)] = j["status"]
    fs, ts = frm.isoformat(), to.isoformat()

    def _dates(mid, fq, lo, hi):
        """This machine's dates for one frequency between lo and hi, after the holiday
        shift and any drag-drop move — exactly the days the calendar would draw."""
        dow = wdays[slot[(mid, fq)]]
        raw = []
        if fq == "weekly":
            d = start + timedelta(days=(dow - start.weekday()) % 7)
            if d < lo - timedelta(days=14):
                d += timedelta(days=7 * ((lo - d).days // 7 - 1))
            while d <= hi:
                if d >= lo - timedelta(days=14):
                    raw.append(d)
                d += timedelta(days=7)
        else:
            nth, moff = phase.get((mid, fq), (0, 0))
            step = FREQ_MONTHS[fq]
            # Counting from the start month, not one interval later. Waiting a whole
            # interval meant the first month of a new programme carried no monthly PM
            # at all, and a yearly machine went twelve months before it was first
            # looked at; the month offset already spreads the long cycles, so a
            # quarterly machine's first visit falls somewhere in the first quarter
            # instead of all of them landing together at the end of it.
            mbase = start.month - 1 + moff
            for k in range(600):
                mm = mbase + step * k
                d = _nth_weekday(start.year + mm // 12, mm % 12 + 1, dow, nth)
                if d is None:
                    continue
                if d > hi:
                    break
                if d >= start and d >= lo - timedelta(days=14):
                    raw.append(d)
        out = []
        for d in raw:
            ds = _next_workday(d, offdays, holidays).isoformat()
            ds = ov.get((mid, ds), ds)                          # apply drag-drop move
            if lo.isoformat() <= ds <= hi.isoformat():
                out.append(ds)
        return out

    # ── work already done must not come round again inside its own cycle ─────────
    #
    # The calendar is recomputed from the checklists every time it is drawn. Upload a
    # new PM plan — a sheet added, a frequency changed, a machine moved to another
    # group — and the route is re-dealt: a machine whose weekly PM was on Monday may now
    # be on Wednesday, a monthly that was the 1st Tuesday may now be the 3rd. The PM
    # the crew already finished on Monday no longer sits on any scheduled day, so the
    # same week showed the machine as due AGAIN on Wednesday.
    #
    # The rule the plant works to: once a PM is done it is not due again until its own
    # cycle comes round — the same Mon–Sun week for weekly, the same calendar month for
    # monthly, the same 3 / 6 / 12-month block (counted from the PM start month) for
    # the long cycles. A PM job that no longer lands on a scheduled day ("orphan") now
    # covers that cycle, and a longer-cycle job covers the shorter ones in it, the same
    # way a quarterly visit already covers the weekly checks on a shared day. Open,
    # still-unfinished PM jobs count too, so the plan never offers a second copy of
    # work that is already on somebody's list. Cancelled and Rejected jobs count for
    # nothing. The finished job stays on the calendar on the day it was really done.
    def _period(fq, ds):
        d = date.fromisoformat(ds)
        if fq == "weekly":
            return ("w", (d - timedelta(days=d.weekday())).isoformat())
        if fq == "monthly":
            return ("m", d.year, d.month)
        step = FREQ_MONTHS.get(fq, 1)
        return (fq, ((d.year * 12 + d.month) - (start.year * 12 + start.month)) // step)

    ext_lo = min(frm, frm - timedelta(days=372))
    # …and just as far AHEAD. Whether a PM is due on a day must not depend on how far
    # the screen asking happens to look. Only jobs up to the last day asked about were
    # read here, so the month calendar — which reaches 25 Sep — saw a PM planned for the
    # 25th and knew 23 Sep's weekly check was covered, while the assign board, asking
    # about 23 Sep alone, did not see it and offered the PM again. The same day read 76
    # on one screen and 78 on the other. A whole cycle either side, always.
    ext_hi = max(to, to + timedelta(days=372))
    pmjobs = {}
    for j in c.execute("SELECT id,machine_id,planned_date,due_date,status,pm_freq,descr,jobtype"
                       " FROM jobs WHERE jobtype='PM' AND machine_id IS NOT NULL"
                       " AND status NOT IN " + VOID_SQL):
        mid = j["machine_id"]
        if mid not in permac:
            continue
        d = (j["planned_date"] or j["due_date"] or "")[:10]
        if not d or not (ext_lo.isoformat() <= d <= ext_hi.isoformat()):
            continue
        try:
            fqj = _job_freq(c, dict(j))
        except Exception:
            fqj = "weekly"
        pmjobs.setdefault(mid, []).append((d, fqj, j["status"]))
    orphans = {}                      # mid -> [(date, rank, freq, status)]
    for mid, lst in pmjobs.items():
        sched = set()
        for fq in permac[mid]["freqs"]:
            if fq in FREQ_WEEKS and (mid, fq) in slot:
                sched.update(_dates(mid, fq, ext_lo, ext_hi + timedelta(days=14)))
        for d, fqj, st in lst:
            if d not in sched:
                orphans.setdefault(mid, []).append(
                    (d, FREQ_ORD.index(fqj) if fqj in FREQ_ORD else 0, fqj, st))

    # One visit per machine per day. Where two sheets still coincide the LONGER one
    # wins — a quarterly service covers the weekly checks, not the other way round.
    best = {}
    for mid, fq in work:
        info = permac[mid]
        rank = FREQ_ORD.index(fq) if fq in FREQ_ORD else 9
        # b438: EVERY PM job of the machine covers its cycle, not only the "orphans".
        # A 6-month PM done on 22 Sep sat on a day the WEEKLY round also uses, so it was
        # not an orphan, did not count — and the 6-month came round again on 6 Oct.
        mine = [(d, FREQ_ORD.index(f) if f in FREQ_ORD else 0, f, st)
                for d, f, st in pmjobs.get(mid, ())]
        for ds in _dates(mid, fq, frm, to):
            if not (fs <= ds <= ts):
                continue
            per = _period(fq, ds)
            if any(r >= rank and od != ds and _period(fq, od) == per for od, r, _f, _s in mine):
                continue                  # already done (or on a list) in this cycle
            k = (ds, mid)
            if k in best and best[k][0] >= rank:
                continue
            # area and climb travel with the row so the calendar can group a day into
            # "Tower B · Weekly · 61" without re-deciding in Javascript what a tower is.
            _ln = info.get("line") or ""
            _ar = _area(_ln)
            best[k] = (rank, {"date": ds, "machine_id": mid, "code": info["code"],
                              "name": info["name"], "line": _ln,
                              "area": _ar, "climb": bool(_ar and CLIMB_RE.search(_ar)),
                              "floor": info.get("floor", ""),
                              "freq": fq, "status": jobstat.get((mid, ds), "Planned")})
    # the orphan jobs themselves stay visible on the day they were actually planned
    for mid, lst in orphans.items():
        info = permac[mid]
        for od, r, fqj, st in lst:
            if not (fs <= od <= ts):
                continue
            k = (od, mid)
            if k in best and best[k][0] >= r:
                continue
            _ln = info.get("line") or ""
            _ar = _area(_ln)
            best[k] = (r, {"date": od, "machine_id": mid, "code": info["code"],
                           "name": info["name"], "line": _ln,
                           "area": _ar, "climb": bool(_ar and CLIMB_RE.search(_ar)),
                           "floor": info.get("floor", ""),
                           "freq": fqj, "status": st})
    return [v[1] for v in best.values()]


def _day_no(start_s):
    if not start_s:
        return 0
    try:
        return (date.fromisoformat(today()) - date.fromisoformat(start_s[:10])).days + 1
    except Exception:
        return 0


@router.get("/config")
async def get_config(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    with closing(db()) as c:
        _ensure(c)
        sd, active = _pm_config(c, _fac(u))
    return {"start_date": sd, "active": bool(active), "day": _day_no(sd)}


@router.post("/start")
async def pm_start(req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        if c.execute("SELECT 1 FROM pm_config WHERE factory_id=?", (fac,)).fetchone():
            c.execute("UPDATE pm_config SET start_date=?,active=1 WHERE factory_id=?", (today(), fac))
        else:
            c.execute("INSERT INTO pm_config(factory_id,start_date,active) VALUES(?,?,1)", (fac, today()))
        c.commit()
    return {"ok": True, "start_date": today()}


@router.post("/stop")
async def pm_stop(req: Request):
    u = require_role(req, "planner", "admin")
    with closing(db()) as c:
        _ensure(c)
        c.execute("UPDATE pm_config SET active=0 WHERE factory_id=?", (_fac(u),))
        c.commit()
    return {"ok": True}


@router.post("/holidays")
async def set_holiday(req: Request):
    """Add or remove a holiday/vacation date (no PM is scheduled on it)."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    dt = (b.get("date") or "")[:10]
    add = bool(b.get("add", True))
    if not dt:
        raise HTTPException(400, "date required")
    with closing(db()) as c:
        _ensure(c)
        r = c.execute("SELECT hours_json FROM factories WHERE id=?", (fac,)).fetchone()
        try:
            cfg = json.loads(r["hours_json"]) if r and r["hours_json"] else {}
        except Exception:
            cfg = {}
        hs = set(cfg.get("holidays", []) or [])
        hs.add(dt) if add else hs.discard(dt)
        cfg["holidays"] = sorted(hs)
        cfg.setdefault("start", cfg.get("start", "07:00"))
        cfg.setdefault("end", cfg.get("end", "21:00"))
        cfg.setdefault("weekly_off", cfg.get("weekly_off", [6]))
        cfg.setdefault("groups", cfg.get("groups", {}))
        c.execute("UPDATE factories SET hours_json=? WHERE id=?", (json.dumps(cfg), fac))
        c.commit()
    return {"ok": True, "holidays": cfg["holidays"]}


@router.get("/cm-calendar")
async def cm_calendar(req: Request):
    """Corrective (CM/BD/IMP/PRJ) jobs by planned date + an unplanned backlog — for the
    CM plan calendar."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    frm = (req.query_params.get("from") or "")[:10]
    to = (req.query_params.get("to") or "")[:10]
    with closing(db()) as c:
        _ensure(c)
        rows = [dict(r) for r in c.execute("""SELECT j.id,j.jobid,j.status,j.jobtype,j.priority,j.planned_date,
                j.descr, j.problem_type, j.fault_category, m.code mcode, m.name mname
              FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id
              WHERE j.jobtype IN ('CM','BD','IMP','PRJ') AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)
                AND j.status NOT IN ('Done','Rejected','Cancelled')""", (fac,))]
        # b417: the Central Electrical CM plan — electrical / other / unknown trade only
        from .elec import is_ce_planner, CeFilter, scope_of
        if is_ce_planner(u) or req.query_params.get("ce") == "1":
            _f = CeFilter(c)
            rows = [r for r in rows if _f.ok(dict(r, factory_id=fac))]
        else:
            _ok = scope_of(c, u)             # b421: a plant planner's CM plan has no electrical CM
            if _ok:
                rows = [r for r in rows if _ok(dict(r, factory_id=fac))]
    days, unplanned, alljobs = {}, [], []
    for r in rows:
        d = (r["planned_date"] or "")[:10]
        # A planner reads a job by how urgent it is and which trade it needs, long before
        # they read its status — so the row has to carry both, not just a status word.
        item = {"id": r["id"], "jobid": r["jobid"], "code": r["mcode"] or r["jobid"],
                "name": r["mname"] or (r["descr"] or ""), "status": r["status"], "jobtype": r["jobtype"],
                "priority": r["priority"], "problem_type": r["problem_type"],
                "fault_category": r["fault_category"], "date": d}
        alljobs.append(item)
        if d:
            if frm <= d <= to:
                days.setdefault(d, []).append(item)
        else:
            unplanned.append(item)
    return {"days": days, "unplanned": unplanned, "all": alljobs, "n_unplanned": len(unplanned)}

# ── Moving a job to another day: what is refused, and why (b379) ───────────────
# 28 Sep, BFLFP: last week's 21 Filling PMs — finished 21-23 Sep and accepted — were
# dragged onto 28 Sep. Nothing stopped it, nothing wrote it down, and because each of
# those machines now "had a PM on the 28th", this week's PM was never offered: the
# calendar showed 21 done and the assign board showed PM 0. Every route that moves a
# job to another day asks this one function first.
FINISHED = ("Done", "ServiceCompleted", "Cancelled", "Rejected")


def move_refusal(c, fac, job, new):
    """None if `job` may be moved to day `new` ('' = take it off its day), otherwise
    the reason, in Thai and English, to show the planner."""
    job = dict(job)
    new = (new or "")[:10]
    jid = job.get("jobid") or job.get("id")
    cur = (job.get("planned_date") or "")[:10]
    if new == cur:
        return None
    if job.get("status") in FINISHED:
        return (f"{jid} ปิดงานแล้ว ({job.get('status')}) — ย้ายวันไม่ได้ /"
                f" {jid} is finished ({job.get('status')}) — a finished job keeps its own day")
    if not new:
        return None
    if new < today():
        return (f"ย้ายงานไปวันที่ผ่านมาแล้ว ({new}) ไม่ได้ /"
                f" a job cannot be moved to a day that has passed ({new})")
    if job.get("jobtype") != "PM":
        return None
    due = (job.get("due_date") or "")[:10]
    if due and new < due:
        return (f"{jid} เป็น PM ของวันที่ {due} — ดึงงาน PM ล่วงหน้ามาทำก่อนกำหนดไม่ได้ /"
                f" {jid} is the PM for {due} — a PM cannot be pulled forward before its own day")
    mid = job.get("machine_id")
    if not mid or not due or new <= due:
        return None
    # Later is allowed — but never onto (or past) the machine's NEXT PM. That day
    # belongs to the next cycle, and a job parked there hides it.
    other = c.execute("""SELECT jobid, COALESCE(NULLIF(planned_date,''),due_date) d FROM jobs
        WHERE jobtype='PM' AND machine_id=? AND id<>? AND status NOT IN ('Cancelled','Rejected')
          AND ((planned_date>? AND planned_date<=?) OR (due_date>? AND due_date<=?))
        ORDER BY d LIMIT 1""", (mid, job["id"], due, new, due, new)).fetchone()
    if other:
        return (f"เครื่องนี้มี PM รอบถัดไป {other['jobid']} วันที่ {other['d']} — ย้าย {jid} ไปถึงวันนั้นไม่ได้ /"
                f" this machine's next PM ({other['jobid']}, {other['d']}) falls on or before {new} — "
                f"{jid} cannot be moved past it")
    try:
        lo = date.fromisoformat(due) + timedelta(days=1)
        hi = date.fromisoformat(new)
        if (hi - lo).days > 400:
            return f"{new}: ไกลเกินไป / too far ahead"
        for o in _occurrences(c, fac, lo, hi):
            if o.get("machine_id") == mid and o["date"] != cur:
                return (f"เครื่องนี้ถึงรอบ PM ถัดไปวันที่ {o['date']} — ย้าย {jid} ไปถึงวันนั้นไม่ได้ /"
                        f" this machine's next PM is due on {o['date']} — {jid} cannot be moved"
                        f" onto or past it")
    except Exception:
        _log.exception("move_refusal: occurrence check failed")
    return None


def log_move(c, job, new, uid):
    """Every move leaves a line in the job's history — who, from what day, to what day."""
    from .chat import log_job_event
    job = dict(job)
    cur = (job.get("planned_date") or "")[:10] or "—"
    who = user_names(c).get(str(uid), "")
    log_job_event(c, job["id"], uid,
                  f"📅 ย้ายวัน {cur} → {new or 'ไม่มีวัน'} โดย {who} /"
                  f" moved from {cur} to {new or 'unplanned'} by {who}")


@router.post("/cm-plan")
async def cm_plan(req: Request):
    """Set (or clear) a corrective job's planned date — used by drag-to-schedule."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    jid = b.get("job_id")
    dt = (b.get("date") or "")[:10]
    lead = b.get("lead_tech")
    if not jid:
        raise HTTPException(400, "job_id required")
    notify = None
    with closing(db()) as c:
        _ensure(c)
        row = c.execute("""SELECT j.id,j.jobid,j.jobtype,j.status,j.machine_id,j.planned_date,j.due_date
            FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id
            WHERE j.id=? AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)""", (jid, fac)).fetchone()
        if not row:
            raise HTTPException(404, "job not found")
        why = move_refusal(c, fac, row, dt)
        if why:
            raise HTTPException(409, why)
        if dt != (row["planned_date"] or "")[:10]:
            log_move(c, row, dt, u["id"])
        if lead:
            c.execute("""UPDATE jobs SET planned_date=?, lead_tech=?,
                status=CASE WHEN status IN ('Reported','WaitingAssignment','WaitingApproval') THEN 'Assigned' ELSE status END
                WHERE id=?""", (dt, lead, jid))
            log_status(c, jid, "Assigned", u["id"])
            set_stage(c, jid)
            r2 = job_row(c, jid)
            notify = (int(lead), r2["jobid"], r2["mcode"] or "")
        else:
            c.execute("UPDATE jobs SET planned_date=? WHERE id=?", (dt, jid))
        c.commit()
    if notify:
        from .push import notify_users
        notify_users([notify[0]], "📋 งานใหม่ถึงคุณ / New job assigned", f"{notify[1]} {notify[2]}", f"/?job={jid}")
    return {"ok": True}


@router.post("/override")
async def override(req: Request):
    """Move a single PM occurrence from one day to another (drag & drop)."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    mid = b.get("machine_id")
    orig = (b.get("orig_date") or "")[:10]
    new = (b.get("new_date") or "")[:10]
    if not (mid and orig and new):
        raise HTTPException(400, "machine_id, orig_date, new_date required")
    with closing(db()) as c:
        _ensure(c)
        if not c.execute("SELECT 1 FROM machines WHERE id=? AND factory_id=?", (mid, fac)).fetchone():
            raise HTTPException(400, "machine not in this factory")
        # A future PM is never pulled forward to an earlier day (b379), and nothing is
        # moved into a day that has already gone. Moving it back to its own day is fine.
        if new != orig and new < orig:
            raise HTTPException(409, f"ดึง PM ของวันที่ {orig} มาทำก่อนกำหนด ({new}) ไม่ได้ /"
                                     f" the PM for {orig} cannot be pulled forward to {new}")
        if new != orig and new < today():
            raise HTTPException(409, f"ย้าย PM ไปวันที่ผ่านมาแล้ว ({new}) ไม่ได้ /"
                                     f" a PM cannot be moved to a day that has passed ({new})")
        c.execute("DELETE FROM pm_overrides WHERE machine_id=? AND orig_date=?", (mid, orig))
        if new != orig:
            c.execute("INSERT INTO pm_overrides(machine_id,orig_date,new_date) VALUES(?,?,?)", (mid, orig, new))
        c.commit()
    return {"ok": True}


@router.get("/planned")
async def planned(req: Request):
    """Planned PM per asset, from the marks, recurring from the program start date
    (Day 1). Weekly = every 7 days; monthly/3M/6M/yearly = every N months from start."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    frm = date.fromisoformat((req.query_params.get("from") or today())[:10])
    to = date.fromisoformat((req.query_params.get("to") or today())[:10])
    elec_only = (req.query_params.get("elec") or "") == "1"
    with closing(db()) as c:
        _ensure(c)
        occ = _occurrences(c, fac, frm, to)
        counts = day_counts(c, fac, frm.isoformat(), to.isoformat(), occ)
        permac = _machine_freqs(c, fac)
        if elec_only:
            # b416: the Central Electrical calendar — the same programme, cut down to the
            # machine-days that carry an electrical (yellow) point, each with how many and
            # who in the CE team has it. Nothing is generated or moved here.
            from .elec import counts as _elc, ELEC_START, count_from as _cf
            _cnt, _m2t = _elc(c, fac), _match_all(c, fac)
            def _ne(mid, fq):
                return _cnt.get((_m2t.get(mid), fq)) or (0, 0)
            occ = [dict(o, n_all=_ne(o["machine_id"], o["freq"])[0], n_el=_ne(o["machine_id"], o["freq"])[1])
                   for o in occ if not o.get("cm") and _ne(o["machine_id"], o["freq"])[1]
                   and o["date"] >= ELEC_START]
            # the job (if raised) and the CE technician, keyed as _make_pm_job keys it
            _jb = {}
            for r in c.execute("""SELECT j.id,j.machine_id,j.planned_date,j.due_date,e.tech,e.started_at,e.done_at
                    FROM jobs j JOIN machines m ON m.id=j.machine_id LEFT JOIN pm_elec e ON e.job_id=j.id
                    WHERE j.jobtype='PM' AND m.factory_id=? AND j.status NOT IN ('Cancelled','Rejected')
                      AND (j.planned_date BETWEEN ? AND ? OR j.due_date BETWEEN ? AND ?)""",
                               (fac, frm.isoformat(), to.isoformat(), frm.isoformat(), to.isoformat())):
                for d in {(r["planned_date"] or "")[:10], (r["due_date"] or "")[:10]}:
                    if d:
                        _jb.setdefault((r["machine_id"], d), dict(r))
            _nm = user_names(c)
            for o in occ:
                j = _jb.get((o["machine_id"], o["date"]))
                o["job_id"] = j["id"] if j else None
                o["ce_tech"] = (j or {}).get("tech")
                o["ce_name"] = _nm.get(str((j or {}).get("tech") or ""), "")
                o["ce_started"] = (j or {}).get("started_at") or ""
                o["ce_done"] = (j or {}).get("done_at") or ""
            _keep = {o["machine_id"] for o in occ}
            permac = {k: v for k, v in permac.items()
                      if any(_ne(k, f)[1] for f in v["freqs"])}
            ce_from = _cf(c)
        sd, active = _pm_config(c, fac)
        offdays, holidays = _cal_config(c, fac)
        cm = [dict(r) for r in c.execute("""SELECT j.id,j.jobid,j.status,j.priority,j.planned_date,
                m.code mcode, m.name mname
              FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id
              WHERE j.jobtype IN ('CM','BD','IMP','PRJ') AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)
                AND j.status NOT IN ('Done','Rejected','Cancelled')
                AND COALESCE(NULLIF(j.planned_date,''),'')<>''""", (fac,))]
    days = {}
    for o in occ:
        days.setdefault(o["date"], []).append(o)
    fss, tss = frm.isoformat(), to.isoformat()
    ncm = 0
    if elec_only:
        cm = []                          # corrective work is not on the electrical calendar
    for j in cm:
        d = (j["planned_date"] or "")[:10]
        if fss <= d <= tss:
            ncm += 1
            days.setdefault(d, []).append({"cm": True, "id": j["id"], "jobid": j["jobid"],
                "code": j["mcode"] or j["jobid"], "name": j["mname"] or "", "status": j["status"], "freq": "cm"})
    # The header strip above the calendar answers "what is this programme?" — how many
    # machines carry a checklist, how much work each frequency asks for in ONE of its
    # own cycles, and how many of the week's days are spent in a tower. Counting that
    # from the visible month would give a different answer every month; it comes from
    # the register instead, so it is the same number a manager quotes.
    from collections import Counter as _C
    _mb = _C()
    for _i in permac.values():
        for _f in _i["freqs"]:
            _mb[_f] += 1
    _climbdays = sorted({d for d, rows in days.items()
                         if any(r.get("climb") for r in rows)
                         and sum(1 for r in rows if r.get("climb")) * 2 >= len(rows)})
    _dows = {date.fromisoformat(d).weekday() for d in _climbdays}
    if elec_only:
        return {"days": days, "counts": {}, "active": bool(active), "start_date": ELEC_START,
                "day": _day_no(ELEC_START), "holidays": sorted(holidays), "weekly_off": sorted(offdays),
                "total": len(occ), "cm": 0, "machines_planned": len(permac),
                "machines_byfreq": {f: sum(1 for k, v in permac.items() if f in v["freqs"] and _ne(k, f)[1])
                                    for f in ("weekly", "monthly", "q3m", "m6", "yearly")},
                "climb_days_per_week": 0, "work_days_per_week": 7 - len(offdays),
                "elec": True, "count_from": ce_from}
    return {"days": days, "counts": counts, "active": bool(active), "start_date": sd, "day": _day_no(sd),
            "holidays": sorted(holidays), "weekly_off": sorted(offdays), "total": len(occ), "cm": ncm,
            "machines_planned": len(permac), "machines_byfreq": dict(_mb),
            "climb_days_per_week": len(_dows), "work_days_per_week": 7 - len(offdays)}


@router.get("/dashboard")
async def pm_dashboard(req: Request):
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    td = date.fromisoformat(today())
    with closing(db()) as c:
        _ensure(c)
        sd, active = _pm_config(c, fac)
        occ = _occurrences(c, fac, td, td + timedelta(days=31))
        permac = _machine_freqs(c, fac)
        planned_machines = len(permac)
        # Assets in the register that no PM template matched — the "Unmatched" tab on
        # the Assets page, counted here so the programme's coverage is one number
        # rather than two screens.
        n_assets = c.execute("SELECT COUNT(*) n FROM machines WHERE factory_id=? AND active=1",
                             (fac,)).fetchone()["n"]
    from collections import Counter
    # How many MACHINES carry each frequency. Different question from "how many PM
    # occurrences fall in the next 30 days", and the one a manager asks first: it is
    # the shape of the programme, not this month's load. A machine counts once per
    # frequency it carries, so the columns add up to more than the machine count.
    mbyfreq = Counter()
    for info in permac.values():
        for f in info["freqs"]:
            mbyfreq[f] += 1
    in_days = lambda n: [o for o in occ if 0 <= (date.fromisoformat(o["date"]) - td).days <= n]
    today_o = [o for o in occ if o["date"] == today()]
    week_o, month_o = in_days(6), in_days(30)
    byfreq = Counter(o["freq"] for o in month_o)
    return {"active": bool(active), "start_date": sd, "day": _day_no(sd),
            "due_today": len(today_o), "due_week": len(week_o), "due_month": len(month_o),
            "machines_today": len({o["machine_id"] for o in today_o}),
            "done_today": len([o for o in today_o if o["status"] == "Done"]),
            "machines_planned": planned_machines,
            "machines_total": n_assets,
            "machines_no_plan": max(0, n_assets - planned_machines),
            "machines_byfreq": {k: mbyfreq.get(k, 0) for k in FREQ_ORD},
            "byfreq": {k: byfreq.get(k, 0) for k in FREQ_ORD}}


@router.get("/calendar")
async def calendar(req: Request):
    """PM work orders positioned by date for the calendar view, plus unplanned and
    missed buckets."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    frm = (req.query_params.get("from") or "")[:10]
    to = (req.query_params.get("to") or "")[:10]
    td = today()
    done = ("Done", "Rejected", "Cancelled")
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT j.id,j.jobid,j.status,j.priority,j.planned_date,j.due_date,j.created_at,
              j.lead_tech, m.code mcode, m.name mname, u.name lead_name
            FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id LEFT JOIN users u ON u.id=j.lead_tech
            WHERE j.jobsource='PM' AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)
              AND j.status NOT IN """ + VOID_SQL, (fac,)).fetchall()
    days, unplanned, missed = {}, [], []
    for r in rows:
        d = dict(r)
        dt = (r["planned_date"] or r["due_date"] or "")[:10]
        item = {"id": r["id"], "jobid": r["jobid"], "status": r["status"], "priority": r["priority"],
                "mcode": r["mcode"], "mname": r["mname"], "lead_name": r["lead_name"], "date": dt}
        if not dt:
            unplanned.append(item)
            continue
        if dt < td and r["status"] not in done:
            missed.append(item)
        if frm <= dt <= to:
            days.setdefault(dt, []).append(item)
    return {"days": days, "unplanned": unplanned, "missed": missed,
            "n_unplanned": len(unplanned), "n_missed": len(missed)}


@router.get("/by-freq")
async def by_freq(req: Request):
    """Machines whose plan for a given frequency falls due within a horizon
    (today / this week / this month), each with that frequency's checklist + images.
    Powers the tree view."""
    u = require_role(req, "planner", "admin", "manager")
    fac = _fac(u)
    freq = req.query_params.get("freq") or "weekly"
    within = req.query_params.get("within") or "month"
    td = date.fromisoformat(today())
    horizon = {"today": 0, "week": 7, "month": 31}.get(within, 3650)
    with closing(db()) as c:
        _ensure(c)
        rows = c.execute("""SELECT ps.interval_days,ps.last_gen,p.start_date,p.template_id,p.machine_id,
              t.machine_type, m.code mcode, m.name mname
            FROM pm_plan_sched ps JOIN pm_plans p ON p.id=ps.plan_id
            JOIN pm_templates t ON t.id=p.template_id LEFT JOIN machines m ON m.id=p.machine_id
            WHERE p.factory_id=? AND p.active=1 AND ps.active=1 AND ps.freq=?
            ORDER BY m.code""", (fac, freq)).fetchall()
        items_cache = {}
        out = []
        for r in rows:
            nd = _next_due(r["interval_days"], r["start_date"], r["last_gen"])
            if not nd:
                continue
            if (date.fromisoformat(nd) - td).days > horizon:
                continue
            tid = r["template_id"]
            if tid not in items_cache:
                items_cache[tid] = [dict(x) for x in c.execute(
                    "SELECT seq,item,normal_status,method,img FROM pm_items WHERE template_id=? AND freq=?"
                    + ACTIVE_ITEM + " ORDER BY seq", (tid, freq))]
            out.append({"machine_id": r["machine_id"], "mcode": r["mcode"], "mname": r["mname"],
                        "machine_type": r["machine_type"], "next_due": nd,
                        "due": nd <= today(), "items": items_cache[tid]})
    return {"freq": freq, "label": FREQ_LABEL.get(freq, freq), "within": within, "machines": out}


@router.patch("/plans/{pid}")
async def edit_plan(pid: int, req: Request):
    require_role(req, "planner", "admin")
    b = await req.json()
    with closing(db()) as c:
        _ensure(c)
        if "active" in b:
            c.execute("UPDATE pm_plans SET active=? WHERE id=?", (1 if b["active"] else 0, pid))
        if "sched" in b:
            s = b["sched"]
            sets = {}
            if "interval_days" in s:
                sets["interval_days"] = max(0, int(s["interval_days"]))
            if "active" in s:
                sets["active"] = 1 if s["active"] else 0
            if sets and s.get("id"):
                c.execute(f"UPDATE pm_plan_sched SET {','.join(k+'=?' for k in sets)} WHERE id=?",
                          (*sets.values(), s["id"]))
        c.commit()
    return {"ok": True}


@router.delete("/plans/{pid}")
async def delete_plan(pid: int, req: Request):
    require_role(req, "planner", "admin")
    with closing(db()) as c:
        _ensure(c)
        c.execute("DELETE FROM pm_plan_sched WHERE plan_id=?", (pid,))
        c.execute("DELETE FROM pm_plans WHERE id=?", (pid,))
        c.commit()
    return {"ok": True}


# ---------- generation ----------

def _make_job(c, template_id, machine_id, machine_type, freq, interval_days, uid, ym):
    items = c.execute("SELECT seq,item,normal_status,method FROM pm_items WHERE template_id=? AND freq=?"
                      + ACTIVE_ITEM + " ORDER BY seq", (template_id, freq)).fetchall()
    lbl = FREQ_LABEL.get(freq, freq)
    lines = [f"PM · {lbl} · {machine_type}  (ทุก {interval_days} วัน / every {interval_days}d)"]
    for it in items:
        lines.append(f"{it['seq']}. {it['item']}  [{it['method'] or 'เช็ค'}]"
                     + (f" — ปกติ: {it['normal_status']}" if it['normal_status'] else ""))
    descr = "\n".join(lines)
    jobid = next_jobid(c, "PM", machine_id=machine_id)
    jid = c.insert_id(
        "INSERT INTO jobs(jobid,jobtype,machine_id,descr,priority,status,jobsource,created_by,created_at)"
        " VALUES(?,?,?,?,?,?,?,?,?)",
        (jobid, "PM", machine_id, descr, 1, "Reported", "PM", uid, now()))
    # the frequency this job is for, so its checklist can ask for those items alone
    try:
        c.execute("UPDATE jobs SET pm_freq=? WHERE id=?", (freq, jid))
    except Exception:
        pass
    return jobid


def _generate(c, uid, fac=None, force_plan=None):
    td = today()
    ym = datetime.now().strftime("%y%m")
    q = ("SELECT ps.*, p.machine_id, p.template_id, p.start_date, p.factory_id, t.machine_type "
         "FROM pm_plan_sched ps JOIN pm_plans p ON p.id=ps.plan_id JOIN pm_templates t ON t.id=p.template_id "
         "WHERE p.active=1 AND ps.active=1")
    args = []
    if fac is not None:
        q += " AND p.factory_id=?"; args.append(fac)
    if force_plan is not None:
        q += " AND p.id=?"; args.append(force_plan)
    created = []
    for s in c.execute(q, args).fetchall():
        nd = _next_due(s["interval_days"], s["start_date"], s["last_gen"])
        if not nd or nd > td:
            continue
        jobid = _make_job(c, s["template_id"], s["machine_id"], s["machine_type"],
                          s["freq"], s["interval_days"], uid, ym)
        c.execute("UPDATE pm_plan_sched SET last_gen=? WHERE id=?", (td, s["id"]))
        if s["machine_id"]:
            c.execute("UPDATE machines SET last_pm_date=? WHERE id=?", (td, s["machine_id"]))
        created.append(jobid)
    return created


def run_due_all():
    """Idempotent daily generation across all factories (called by the scheduler)."""
    with closing(db()) as c:
        _ensure(c)
        sys_uid = (c.execute("SELECT id FROM users WHERE role IN ('admin','planner') ORDER BY id LIMIT 1").fetchone() or {"id": 1})["id"]
        created = _generate(c, sys_uid, fac=None)
        c.commit()
    return created


@router.post("/run-due")
async def run_due(req: Request):
    u = require_role(req, "planner", "admin")
    with closing(db()) as c:
        _ensure(c)
        created = _generate(c, u["id"], fac=_fac(u))
        c.commit()
    return {"ok": True, "created": len(created), "jobids": created}


def _assets_or_400(c, fac):
    """A PM plan needs machines to hang on. Refusing an empty factory is also what
    stops a plan being loaded into the wrong plant by accident."""
    n = c.execute("SELECT COUNT(*) FROM machines WHERE factory_id=?", (fac,)).fetchone()[0]
    if not n:
        raise HTTPException(400, "โรงงานนี้ยังไม่มีทะเบียนสินทรัพย์ — นำเข้าสินทรัพย์ก่อน "
                                 "แล้วจึงอัปโหลดแผน PM / this factory has no assets yet — "
                                 "import the asset register first, then upload the PM plan")
    return n


def _apply_plan(c, fac, plan, keep_plans, who=None):
    """Replace this factory's checklists with `plan`. Caller commits."""
    form = plan.get("form_no", "")
    # An import gives every template a brand-new id. Anything that points at a
    # template by id — a saved PM plan, and an asset pinned to a sheet by hand on
    # the Assets page — has to be carried across by NAME, or it is left pointing
    # at a row that no longer exists.
    oldname = {r["id"]: (r["machine_type"] or "").strip()
               for r in c.execute("SELECT id,machine_type FROM pm_templates WHERE factory_id=?", (fac,))}
    pinned = [(r["machine_id"], oldname.get(r["template_id"], ""))
              for r in c.execute("SELECT machine_id,template_id FROM pm_members WHERE factory_id=?", (fac,))]
    kept_plans = [(r["id"], oldname.get(r["template_id"], ""))
                  for r in c.execute("SELECT id,template_id FROM pm_plans WHERE factory_id=?", (fac,))]
    old = list(oldname)
    if old and not keep_plans:
        ph = ",".join("?" * len(old))
        pids = [r["id"] for r in c.execute(f"SELECT id FROM pm_plans WHERE template_id IN ({ph})", old)]
        if pids:
            ph2 = ",".join("?" * len(pids))
            c.execute(f"DELETE FROM pm_plan_sched WHERE plan_id IN ({ph2})", pids)
            c.execute(f"DELETE FROM pm_plans WHERE id IN ({ph2})", pids)
    # ── items are MATCHED, not wiped ──────────────────────────────────────────────
    # This used to delete every template and item and insert the file afresh, so every
    # item got a new id and every saved OK/NG answer lost the item it belonged to. Now
    # a sheet is kept by NAME and an item by its wording (then by its number): same
    # wording = the same item, untouched; a different wording at the same number = a
    # changed item, same id; anything new is added; anything the file no longer has is
    # hidden (active=0), never deleted.
    _ensure_records(c)
    byname = {}
    for r in c.execute("SELECT id,machine_type FROM pm_templates WHERE factory_id=?", (fac,)):
        byname.setdefault((r["machine_type"] or "").strip(), r["id"])
    st = {"kept": 0, "changed": 0, "added": 0, "removed": 0}
    overwritten = []
    norm = lambda x: re.sub(r"\s+", " ", str(x or "")).strip()
    n_t = n_i = 0
    newid, stated = {}, {}
    for t in plan.get("templates", []):
        name = (t["machine_type"] or "").strip()
        if name in byname and name not in newid:
            tid = byname[name]
            c.execute("UPDATE pm_templates SET form_no=?, color=? WHERE id=?",
                      (form, t.get("color", ""), tid))
        else:
            tid = c.insert_id("INSERT INTO pm_templates(factory_id,machine_type,form_no,color) VALUES(?,?,?,?)",
                              (fac, t["machine_type"], form, t.get("color", "")))
        newid[name] = tid
        # A master workbook states which machines each sheet covers. The BFLFP plan does
        # not, so the pins below are re-linked by template NAME — a guess that has to be
        # made when the document is silent. Where the document speaks, it wins.
        if t.get("members"):
            stated[tid] = list(t["members"])
        n_t += 1
        olds = [dict(r) for r in c.execute("SELECT * FROM pm_items WHERE template_id=? ORDER BY seq, id", (tid,))]
        free = {o["id"]: o for o in olds}
        for it in t.get("items", []):
            want = {"seq": it.get("seq"), "item": it.get("item", "") or "",
                    "normal_status": it.get("normal", "") or "", "method": it.get("method", "") or "",
                    "freq": it.get("freq", "weekly") or "weekly"}
            m = next((o for o in free.values() if (o["active"] if o["active"] is not None else 1)
                      and norm(o["item"]) == norm(want["item"])), None)
            if m is None:
                m = next((o for o in free.values() if not (o["active"] if o["active"] is not None else 1)
                          and norm(o["item"]) == norm(want["item"])), None)
            if m is None:
                m = next((o for o in free.values() if (o["active"] if o["active"] is not None else 1)
                          and o["seq"] == want["seq"]), None)
            n_i += 1
            if m is None:
                _nid = c.insert_id("INSERT INTO pm_items(template_id,seq,item,normal_status,method,freq,img,active,app_edit)"
                          " VALUES(?,?,?,?,?,?,?,1,0)",
                          (tid, want["seq"], want["item"], want["normal_status"], want["method"],
                           want["freq"], it.get("img", "") or ""))
                if "elec" in it:                       # b413: yellow cell = electrical point
                    c.execute("UPDATE pm_items SET elec=? WHERE id=?", (1 if it["elec"] else 0, _nid))
                st["added"] += 1
                continue
            del free[m["id"]]
            was_active = (m["active"] if m["active"] is not None else 1)
            differs = any(norm(m.get(k)) != norm(want[k]) for k in ITEM_FIELDS)
            if not was_active:
                st["added"] += 1
                if m.get("app_edit"):
                    overwritten.append(f"{name} #{want['seq']}")   # removed in the app, back from the file
            elif differs:
                st["changed"] += 1
                if m.get("app_edit"):
                    overwritten.append(f"{name} #{want['seq']}")
            else:
                st["kept"] += 1
            c.execute("UPDATE pm_items SET seq=?,item=?,normal_status=?,method=?,freq=?,img=?,active=1,app_edit=0"
                      " WHERE id=?", (want["seq"], want["item"], want["normal_status"], want["method"],
                                      want["freq"], it.get("img") or m.get("img") or "", m["id"]))
            # b413: the master marks electrical points in yellow. Only a document that says
            # (the master reader always does) sets it; the BFLFP plan leaves it alone.
            if "elec" in it:
                c.execute("UPDATE pm_items SET elec=? WHERE id=?", (1 if it["elec"] else 0, m["id"]))
        for o in free.values():
            if o["active"] if o["active"] is not None else 1:
                st["removed"] += 1
                if o.get("app_edit"):
                    overwritten.append(f"{name} #{o['seq']}")
                c.execute("UPDATE pm_items SET active=0, app_edit=0 WHERE id=?", (o["id"],))
    # sheets the file no longer has: the sheet goes, its items are hidden, not deleted
    for name, tid in byname.items():
        if name in newid:
            continue
        n_gone = c.execute("SELECT COUNT(*) n FROM pm_items WHERE template_id=?" + ACTIVE_ITEM,
                           (tid,)).fetchone()["n"]
        st["removed"] += n_gone
        c.execute("UPDATE pm_items SET active=0, app_edit=0 WHERE template_id=?", (tid,))
        c.execute("DELETE FROM pm_templates WHERE id=?", (tid,))
    answers = c.execute(
        "SELECT COUNT(*) n FROM pm_results r JOIN jobs j ON j.id=r.job_id"
        " JOIN machines m ON m.id=j.machine_id WHERE m.factory_id=?", (fac,)).fetchone()["n"]
    _item_log(c, fac, None, None, "upload", "",
              json.dumps({**st, "templates": n_t, "overwritten": overwritten}, ensure_ascii=False),
              who)
    # re-link the pins and the saved plans to the template of the same name;
    # anything whose sheet is gone from the workbook is dropped rather than left dangling
    c.execute("DELETE FROM pm_members WHERE factory_id=?", (fac,))
    n_pin = n_stated = 0
    missing = []
    if stated:
        # bind by asset code, in one lookup rather than one query per machine
        codes = {cd for v in stated.values() for cd in v}
        idof = {}
        cl = sorted(codes)
        for i in range(0, len(cl), 400):
            chunk = cl[i:i + 400]
            ph = ",".join("?" * len(chunk))
            for r in c.execute(f"SELECT id,code FROM machines WHERE factory_id=? AND code IN ({ph})",
                               (fac, *chunk)):
                idof[r["code"]] = r["id"]
        seen = set()
        for tid, cds in stated.items():
            for cd in cds:
                mid_ = idof.get(cd)
                if mid_ is None:
                    missing.append(cd)          # in the workbook, not in the register
                    continue
                if mid_ in seen:
                    continue                    # one machine, one checklist
                seen.add(mid_)
                c.execute("INSERT INTO pm_members(template_id,machine_id,factory_id) VALUES(?,?,?)",
                          (tid, mid_, fac))
                n_stated += 1
    for mid_, name in pinned:
        if stated:
            break                               # the document said it; do not guess over it
        if name and name in newid:
            c.execute("INSERT INTO pm_members(template_id,machine_id,factory_id) VALUES(?,?,?)",
                      (newid[name], mid_, fac))
            n_pin += 1
    n_plan = 0
    for pid, name in kept_plans:
        if name and name in newid:
            c.execute("UPDATE pm_plans SET template_id=? WHERE id=?", (newid[name], pid))
            n_plan += 1
        elif keep_plans:
            c.execute("DELETE FROM pm_plan_sched WHERE plan_id=?", (pid,))
            c.execute("DELETE FROM pm_plans WHERE id=?", (pid,))
    for g in plan.get("color_groups", []):
        if g.get("color") and not c.execute(
                "SELECT 1 FROM pm_groups WHERE factory_id=? AND color=?", (fac, g["color"])).fetchone():
            c.execute("INSERT INTO pm_groups(factory_id,color,name) VALUES(?,?,?)",
                      (fac, g["color"], g.get("name") or g["color"]))
    _MATCH_CACHE.pop(fac, None)
    return {"ok": True, "templates": n_t, "items": n_i, "pins": n_pin + n_stated,
            "items_kept": st["kept"], "items_changed": st["changed"], "items_added": st["added"],
            "items_removed": st["removed"], "app_edits_overwritten": overwritten,
            "answers_kept": answers,
            "bound_from_file": n_stated, "plans": n_plan,
            "not_in_register": sorted(set(missing))[:50],
            "not_in_register_n": len(set(missing))}


@router.post("/import")
async def import_plan(req: Request):
    """Reload the checklists from the copy already on the server (data/pm_plan.json)."""
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    keep_plans = bool((await _safe_json(req)).get("keep_plans"))
    try:
        plan = json.load(open(PLAN_FILE, encoding="utf-8"))
    except Exception as e:
        raise HTTPException(400, f"pm_plan.json not found or unreadable: {e}")
    with closing(db()) as c:
        _ensure(c)
        n_ast = _assets_or_400(c, fac)
        res = _apply_plan(c, fac, plan, keep_plans, who=u)
        c.commit()
    res["assets"] = n_ast
    return res


@router.post("/import-xlsx")
async def import_plan_upload(req: Request):
    """Upload the PM Plan workbook itself (base64 in JSON) and load it into this
    factory: the asset register has to exist first, then the sheets replace whatever
    checklists are here, and the part photos are written into static/pmimg."""
    import base64
    import os as _os
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    b = await req.json()
    data = b.get("b64") or ""
    if "," in data[:80]:                       # tolerate a data:...;base64, prefix
        data = data.split(",", 1)[1]
    try:
        raw = base64.b64decode(data)
    except Exception:
        raise HTTPException(400, "bad file data")
    if not raw:
        raise HTTPException(400, "empty file")
    # the two uploads look alike in the file picker, so say plainly when the wrong
    # one lands here rather than complaining that the factory is empty
    fname = (b.get("name") or "")
    tag = fname.upper().replace(" ", "")
    if "ASSET" in tag and "REVIEW" in tag:
        raise HTTPException(400, "นี่คือไฟล์ทะเบียนสินทรัพย์ ไม่ใช่แผน PM — ใช้เมนู “นำเข้าสินทรัพย์” "
                                 "ทางซ้ายก่อน / this is the asset register, not the PM plan — "
                                 "use “Import assets” in the left menu first")
    want = "PC" if "BFLPC" in tag else "FP" if "BFLFP" in tag else None
    if want and u.get("factory_code") and want != u["factory_code"]:
        raise HTTPException(400, f"ไฟล์นี้เป็นของโรงงาน {want} — คุณเข้าสู่ระบบโรงงาน {u.get('factory_code')} "
                                 f"(this file is for {want}; you are logged into {u.get('factory_code')})")
    from .config import STATIC
    # Which document is this? Two shapes are in use and they share nothing but the word
    # "PM": the BFLFP PM Plan (one sheet per machine type, frequency in the fill colour
    # of the ลำดับ cell, a part photo per row) and the machine-register master (a sheet
    # per machine GROUP, frequency in tick columns, method in its own column, no photos).
    # Decided by looking inside the file, not by its name — both get renamed constantly.
    from . import pm_import_master as _mst
    from .pm_import import plan_from_workbook
    with closing(db()) as c:
        _ensure(c)
        n_ast = _assets_or_400(c, fac)         # checked before a single row is touched
        try:
            _wb = _mst._load(raw)
            _is_master = _mst.looks_like_master(_wb)
            _wb.close()
        except ValueError:
            _is_master = False
        mlist = []
        try:
            if _is_master:
                plan, mlist = _mst.plan_from_workbook(raw, fac)
                images = {}                     # this document carries no part photos
            else:
                plan, images = plan_from_workbook(raw, fac)
        except ValueError as e:
            raise HTTPException(400, str(e))
        # The filename guard above catches "BFLPC" or "BFLFP" in the name. A
        # machine-register master is named after its FORM — "SD-SP-ENG02-01 … master
        # 2026" — and names no plant at all, so that guard can never fire for one and a
        # planner in the wrong plant would replace its checklists with another plant's.
        # The workbook's own asset codes settle it: ask which plant they actually belong
        # to, and refuse if it is not this one.
        if _is_master and mlist:
            codes = [m["code"] for m in mlist if m.get("code")]
            qs = ",".join("?" * len(codes))
            owner = {}
            for r in c.execute(
                    "SELECT m.factory_id fid, f.code fcode, COUNT(*) n FROM machines m"
                    " LEFT JOIN factories f ON f.id=m.factory_id"
                    " WHERE m.code IN (%s) GROUP BY m.factory_id" % qs, codes):
                owner[r["fid"]] = (r["fcode"] or "?", r["n"])
            here = owner.get(fac, ("", 0))[1]
            best = max(owner.items(), key=lambda kv: kv[1][1], default=None)
            # Only refuse when the register can actually answer: a plant whose codes are
            # nearly all somewhere else is a mix-up, one that simply has few of them yet
            # is a first import and must still be allowed.
            if best and best[0] != fac and best[1][1] >= 20 and here < best[1][1] / 4:
                raise HTTPException(400,
                    "ไฟล์นี้เป็นทะเบียนของโรงงาน %s (รหัสเครื่อง %d รายการตรงกับโรงงานนั้น, "
                    "ตรงกับโรงงานนี้ %d) — คุณเข้าสู่ระบบโรงงาน %s / this workbook is %s's "
                    "register (%d of its asset codes belong to %s, %d to this plant) — you "
                    "are logged into %s"
                    % (best[1][0], best[1][1], here, u.get("factory_code") or "?",
                       best[1][0], best[1][1], best[1][0], here, u.get("factory_code") or "?"))
        if any((t["machine_type"] or "").strip().lower() == "review" for t in plan["templates"]):
            raise HTTPException(400, "ไฟล์นี้ดูเหมือนทะเบียนสินทรัพย์ ไม่ใช่แผน PM / "
                                     "this looks like the asset register, not a PM plan")
        res = _apply_plan(c, fac, plan, bool(b.get("keep_plans")), who=u)
        # Checklists alone generate nothing. A plant with no PM start date has an empty
        # PM screen and no explanation for it, so the upload reports whether the
        # programme is running and the page can say what to press next.
        _cfg = c.execute("SELECT start_date,active FROM pm_config WHERE factory_id=?",
                         (fac,)).fetchone()
        res["pm_started"] = bool(_cfg and _cfg["active"])
        res["pm_start_date"] = (_cfg["start_date"] if _cfg else "") or ""
        c.commit()
    # the checklists are in; the part photos only reach a phone through static/pmimg
    imgd = _os.path.join(STATIC, "pmimg")
    _os.makedirs(imgd, exist_ok=True)
    written = 0
    for name, blob in images.items():
        try:
            with open(_os.path.join(imgd, name), "wb") as fh:
                fh.write(blob)
            written += 1
        except Exception:
            pass
    # keep the server-side copy in step, so the plain Import button reloads the same thing
    try:
        with open(PLAN_FILE, "w", encoding="utf-8") as fh:
            json.dump(plan, fh, ensure_ascii=False, indent=1)
    except Exception:
        pass
    res.update({"assets": n_ast, "photos": written, "file": fname,
                "layout": "master" if _is_master else "bflfp"})
    if _is_master:
        pend = next((t for t in plan["templates"]
                     if t["machine_type"] == _mst.PENDING), None)
        # rows the reader refused to guess at — handed back so somebody can fix the file
        notes = plan.get("notes") or {}
        res.update({"mlist": len(mlist),
                    "pending": len(pend["members"]) if pend else 0,
                    "misplaced_rows": notes.get("misplaced") or [],
                    "listed_twice": notes.get("doubled") or []})
    return res


async def _safe_json(req):
    try:
        return await req.json()
    except Exception:
        return {}


@router.post("/clear-jobs")
async def clear_jobs(req: Request):
    u = require_role(req, "planner", "admin")
    fac = _fac(u)
    with closing(db()) as c:
        _ensure(c)
        ids = [r["id"] for r in c.execute(
            "SELECT j.id FROM jobs j LEFT JOIN machines m ON m.id=j.machine_id "
            "WHERE j.jobsource='PM' AND (COALESCE(m.factory_id,j.factory_id)=? OR COALESCE(m.factory_id,j.factory_id) IS NULL)", (fac,))]
        if ids:
            ph = ",".join("?" * len(ids))
            c.execute(f"DELETE FROM job_events WHERE job_id IN ({ph})", ids)
            try:
                c.execute(f"DELETE FROM timelogs WHERE job_id IN ({ph})", ids)
            except Exception:
                pass
            c.execute(f"DELETE FROM jobs WHERE id IN ({ph})", ids)
        c.execute("UPDATE pm_plan_sched SET last_gen='' WHERE plan_id IN "
                  "(SELECT id FROM pm_plans WHERE factory_id=?)", (fac,))
        c.commit()
    return {"ok": True, "deleted": len(ids)}


# ── the checklist a technician actually fills in ────────────────────────────────
# The printed form carries every frequency at once — five Week boxes, 1M, 3M, 6M, 1Y —
# because one sheet of paper had to serve the whole year. A job does not: it is the
# monthly PM on this machine, so it asks for the monthly items and nothing else. What
# comes back is OK or NG per item, a reason on every NG, and the name of the person who
# looked. That is the record; the frequency grid was scaffolding for paper.

RESULT_OK, RESULT_NG = "OK", "NG"


def _job_freq(c, job):
    """Which frequency of the sheet this PM job is for."""
    f = (job.get("pm_freq") or "").strip()
    if f in FREQ_STEP or f == "each_op":
        return f
    # jobs created before the column existed: the label is the first line of the
    # description, written by _make_pm_job as "PM · <label> · <template>"
    head = str(job.get("descr") or "").split("\n")[0]
    for k, lbl in FREQ_LABEL.items():
        if f"· {lbl} ·" in head or head.endswith("· " + lbl):
            return k
    return "weekly"


def _live_items(c, job):
    """(freq, [items]) the sheet holds TODAY for this PM job's machine and frequency."""
    fac = c.execute("SELECT factory_id FROM machines WHERE id=?",
                    (job["machine_id"],)).fetchone()
    fac = fac["factory_id"] if fac else 2
    freq = _job_freq(c, job)
    tid = _match_all(c, fac).get(job["machine_id"])
    if not tid:
        return freq, []
    return freq, [dict(r) for r in c.execute(
        "SELECT id,seq,item,normal_status,method,img,freq FROM pm_items"
        " WHERE template_id=? AND freq=?" + ACTIVE_ITEM + " ORDER BY seq", (tid, freq))]


def snapshot_job_items(c, job):
    """Freeze the list a PM job is worked from. Called when the job is started and when
    the first answer is saved; a job that already has a frozen list keeps it. Returns
    the number of rows written. Caller commits."""
    if (job.get("jobtype") or "") != "PM" or not job.get("machine_id"):
        return 0
    _ensure_records(c)
    if c.execute("SELECT 1 FROM pm_job_items WHERE job_id=? LIMIT 1", (job["id"],)).fetchone():
        return 0
    freq, rows = _live_items(c, job)
    if not rows:
        # the sheet no longer has this frequency: rebuild from the stored answers
        rows = [{"id": r["item_id"], "seq": r["seq"], "item": r["item_text"] or "",
                 "normal_status": r["item_normal"] or "", "method": r["item_method"] or "",
                 "img": "", "freq": r["item_freq"] or freq}
                for r in c.execute("SELECT * FROM pm_results WHERE job_id=? AND COALESCE(item_text,'')<>''"
                                   " ORDER BY seq", (job["id"],))]
    for r in rows:
        c.execute("INSERT INTO pm_job_items(job_id,item_id,seq,item,normal_status,method,img,freq)"
                  " VALUES(?,?,?,?,?,?,?,?)",
                  (job["id"], r["id"], r["seq"], r["item"] or "", r["normal_status"] or "",
                   r["method"] or "", r.get("img") or "", r.get("freq") or freq))
    return len(rows)


def job_checklist(c, job):
    """{freq, label, items:[...], done, total, ng} for a PM job, else None.

    Every item carries whatever has already been answered, so a technician who puts the
    phone down mid-round picks up where they left off rather than starting again.

    Which list: a job that has been started keeps the list it started with
    (pm_job_items); one not started yet follows the sheet as it is now, so an edit
    reaches every PM that has not begun. Finished jobs are always read from their
    frozen list, and each answer carries its own wording besides, so no later edit or
    Excel upload can change what a finished PM record says.
    """
    if (job.get("jobtype") or "") != "PM" or not job.get("machine_id"):
        return None
    try:
        _ensure(c)
        freq = _job_freq(c, job)
        rows = [dict(r) for r in c.execute(
            "SELECT item_id id,seq,item,normal_status,method,img FROM pm_job_items"
            " WHERE job_id=? ORDER BY seq, id", (job["id"],))]
        frozen = bool(rows)
        if not rows:
            freq, rows = _live_items(c, job)
        got = {r["item_id"]: dict(r) for r in c.execute(
            "SELECT item_id,result,remark,tech_name,at,seq,item_text,item_normal,item_method"
            " FROM pm_results WHERE job_id=?", (job["id"],))}
        if not rows:
            # no frozen list and nothing on the sheet: whatever was answered is the record
            rows = [{"id": a["item_id"], "seq": a["seq"], "item": a["item_text"] or "",
                     "normal_status": a["item_normal"] or "", "method": a["item_method"] or "",
                     "img": ""} for a in sorted(got.values(), key=lambda x: x["seq"] or 0)
                    if (a["item_text"] or "")]
        if not rows:
            return None
        out = []
        for r in rows:
            a = got.get(r["id"]) or {}
            out.append({"item_id": r["id"], "seq": r["seq"],
                        # an answer states what it was answered against
                        "item": (a.get("item_text") or r["item"] or ""),
                        "normal": (a.get("item_normal") or r["normal_status"] or ""),
                        "method": (a.get("item_method") or r["method"] or ""),
                        "img": r["img"] or "",
                        "result": a.get("result") or "", "remark": a.get("remark") or "",
                        "tech_name": a.get("tech_name") or "", "at": a.get("at") or ""})
        # b413: which points are electrical — the Central Engineering team's, not the
        # plant technician's. Read live from the item, so a sheet re-uploaded with its
        # yellow marks applies to jobs already raised.
        try:
            from .elec import elec_items as _elx
            ELX = _elx(c, [i["item_id"] for i in out])
        except Exception:
            ELX = set()
        for i in out:
            i["elec"] = 1 if i["item_id"] in ELX else 0
        done = sum(1 for i in out if i["result"] in (RESULT_OK, RESULT_NG))
        ng = [i for i in out if i["result"] == RESULT_NG]
        gen = [i for i in out if not i["elec"]]
        el = [i for i in out if i["elec"]]
        _ans = lambda i: i["result"] in (RESULT_OK, RESULT_NG)        # noqa: E731
        _nor = lambda i: i["result"] == RESULT_NG and not (i["remark"] or "").strip()  # noqa: E731
        res = {"freq": freq, "label": FREQ_LABEL.get(freq, freq), "items": out,
               "total": len(out), "done": done, "ng": len(ng), "frozen": frozen,
               # what still stands between the PLANT crew and the Stop button — the
               # electrical points are not theirs to answer, so they do not block it
               "missing": [i["seq"] for i in gen if not _ans(i)],
               "ng_no_remark": [i["seq"] for i in gen if _nor(i)],
               "el_total": len(el), "el_done": sum(1 for i in el if _ans(i)),
               "el_missing": [i["seq"] for i in el if not _ans(i)],
               "el_ng_no_remark": [i["seq"] for i in el if _nor(i)]}
        res["el_missing_n"] = len(res["el_missing"])
        if el:
            try:
                from .elec import state_of as _elst
                res["elec_state"] = _elst(c, job["id"])
            except Exception:
                res["elec_state"] = None
        return res
    except Exception as e:
        _log.warning("[pm] checklist for job %s failed: %s", job.get("id"), e)
        return None


def checklist_block(c, job):
    """Why a PM job may not be closed yet — empty string when it may."""
    cl = job_checklist(c, job)
    if not cl:
        return ""
    if cl["missing"]:
        return ("ยังตรวจไม่ครบ %d รายการ (ข้อ %s) / %d item(s) not answered yet (no. %s)"
                % (len(cl["missing"]), ", ".join(map(str, cl["missing"][:8])),
                   len(cl["missing"]), ", ".join(map(str, cl["missing"][:8]))))
    if cl["ng_no_remark"]:
        return ("ข้อที่ไม่ผ่าน (NG) ต้องใส่หมายเหตุ — ข้อ %s / an NG needs a remark — no. %s"
                % (", ".join(map(str, cl["ng_no_remark"])),
                   ", ".join(map(str, cl["ng_no_remark"]))))
    return ""


@router.get("/job/{jid}/checklist")
async def get_job_checklist(jid: int, req: Request):
    require_role(req, "technician", "planner", "admin", "manager", "operator", "engcenter")
    with closing(db()) as c:
        cl = job_checklist(c, job_row(c, jid))
    if cl is None:
        raise HTTPException(404, "no PM checklist for this job")
    return cl


@router.post("/job/{jid}/checklist")
async def save_job_checklist(jid: int, req: Request):
    """Answer one or more items. Saved as they are tapped, not held to the end — a
    phone that dies halfway up a ladder must not cost the round."""
    u = require_role(req, "technician", "planner", "admin")
    b = await req.json()
    items = b.get("items") or []
    if not isinstance(items, list):
        raise HTTPException(400, "items must be a list")
    who = b.get("as_tech")
    with closing(db()) as c:
        _ensure(c)
        job = job_row(c, jid)
        if (job.get("jobtype") or "") != "PM":
            raise HTTPException(400, "not a PM job")
        names = user_names(c)
        tech = int(who) if str(who or "").isdigit() else u["id"]
        tname = names.get(str(tech), u.get("name") or "")
        # b413: who may answer which point. A Central Electrical technician answers the
        # electrical (yellow) points and nothing else; a plant technician everything but.
        # Planner and admin enter on somebody's behalf and may answer either.
        ce_side = None
        if u["role"] == "technician":
            ce_side = (u.get("team") or "") == "CE"
            if ce_side:
                from .elec import on_crew as _onc
                if _onc(job, u):
                    ce_side = None          # b428: on the plant crew — the whole sheet is his
        try:
            from .elec import elec_items as _elx
            ELX = _elx(c, [it.get("item_id") for it in items if str(it.get("item_id") or "").isdigit()])
        except Exception:
            ELX = set()
        if ce_side is not None:
            bad = [it for it in items if str(it.get("item_id") or "").isdigit()
                   and ((int(it["item_id"]) in ELX) != ce_side)]
            if bad:
                raise HTTPException(403, ("ข้อนี้เป็นงานของช่างโรงงาน / this point is the plant technicians'"
                                          if ce_side else
                                          "ข้อไฟฟ้า — ช่างไฟฟ้าส่วนกลาง (CE) เป็นผู้ตรวจ / an electrical point —"
                                          " answered by Central Electrical"))
            if ce_side:
                st = c.execute("SELECT started_at FROM pm_elec WHERE job_id=?", (jid,)).fetchone()
                if not st or not (st["started_at"] or ""):
                    raise HTTPException(409, "กด “เริ่มตรวจ” ก่อน / press Start check first")
        ts = now()
        # the first answer freezes the list, if Start did not already
        snapshot_job_items(c, job)
        known = {r["item_id"]: dict(r) for r in c.execute(
            "SELECT item_id,seq,item,normal_status,method,freq FROM pm_job_items WHERE job_id=?", (jid,))}
        for it in items:
            try:
                iid = int(it.get("item_id"))
            except (TypeError, ValueError):
                continue
            res = (it.get("result") or "").upper()
            if res not in (RESULT_OK, RESULT_NG, ""):
                continue
            rem = str(it.get("remark") or "")[:400]
            seq = int(it.get("seq") or 0)
            src = known.get(iid)
            if src is None:
                r1 = c.execute("SELECT id item_id,seq,item,normal_status,method,freq FROM pm_items"
                               " WHERE id=?", (iid,)).fetchone()
                src = dict(r1) if r1 else {}
            if known and iid not in known:
                continue                  # not on the list this job is worked from
            txt, nrm = src.get("item") or "", src.get("normal_status") or ""
            mth, frq = src.get("method") or "", src.get("freq") or ""
            ex = c.execute("SELECT id FROM pm_results WHERE job_id=? AND item_id=?",
                           (jid, iid)).fetchone()
            if ex:
                c.execute("UPDATE pm_results SET result=?,remark=?,seq=?,tech=?,tech_name=?,at=?,"
                          " item_text=?,item_normal=?,item_method=?,item_freq=? WHERE id=?",
                          (res, rem, seq, tech, tname, ts, txt, nrm, mth, frq, ex["id"]))
            else:
                c.execute("INSERT INTO pm_results(job_id,item_id,seq,result,remark,tech,tech_name,at,"
                          "item_text,item_normal,item_method,item_freq) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                          (jid, iid, seq, res, rem, tech, tname, ts, txt, nrm, mth, frq))
        c.commit()
        cl = job_checklist(c, job_row(c, jid))
    return {"ok": True, "checklist": cl}
