"""Shared asset-register import — used by the Planner UI upload endpoint.

Two workbook shapes are accepted, because the register is edited in both:

  * the **Asset-Review template** — sheet `Review`, fixed columns, data from row 6.
    This is the form handed out for a first load, and its layout is read positionally.

  * whatever **Export Excel** produced — sheet `Assets`, a header row, data under it.
    Export → edit in Excel → import back is the loop a planner actually uses when
    filling in a column for the whole plant (locations, serial numbers, sizes), and
    it used to fail with "no asset rows found" because the exporter and the importer
    disagreed about which row and column anything was in. The export layout is read
    by its **header names**, so a reordered or partial sheet still lands correctly and
    only the columns present are written.

UPSERTS into the machines table for ONE factory, keyed on the asset code. Mirrors
import_assets.py (the CLI script) but runs through the portable db() connection so it
works on SQLite and PostgreSQL alike.
"""
import io, datetime

# machines field  <-  Excel column (0-based index; B=1, C=2, ...) — same as import_assets.py
MAP = [("code", 1), ("name", 2), ("category", 3), ("asset_group", 4), ("kpi_class", 5),
       ("criticality", 6), ("line", 7), ("floor", 8), ("department", 9), ("brand_model", 10),
       ("manufacturer", 11), ("size", 12), ("serial_no", 13), ("year_install", 14),
       ("last_pm_date", 15), ("remark", 16)]
FIELDS = [m[0] for m in MAP]

# Header text (lower-cased, trimmed) -> machines field, for the export-shaped sheet.
# "PM group" and "PM sheet" are deliberately absent: both are worked out by the PM
# matcher from the machine's name, so importing them would write a guess back over the
# thing that produced it.
HEADERS = {
    "asset id": "code", "code": "code", "asset code": "code", "รหัสเครื่อง": "code",
    "name": "name", "ชื่อเครื่อง": "name",
    "category": "category",
    "group": "asset_group", "asset group": "asset_group",
    "criticality": "criticality", "crit": "criticality",
    "line / location": "line", "line/location": "line", "location": "line",
    "line": "line", "สถานที่": "line",
    "floor": "floor", "ชั้น": "floor",
    "department": "department",
    "manufacturer": "manufacturer",
    "brand / model": "brand_model", "brand/model": "brand_model", "model": "brand_model",
    "serial no": "serial_no", "serial no.": "serial_no", "serial": "serial_no",
    "size": "size",
    "year installed": "year_install", "year install": "year_install",
    "last pm": "last_pm_date",
    "pm every (days)": "pm_freq_days", "pm every": "pm_freq_days",
    "active": "active",
    "remark": "remark",
}


def _header_row(ws, limit=8):
    """Find the header row and the field -> column index map it describes.

    Returns (row_number, {field: col_index}) or (None, None) when this does not look
    like a header-driven sheet at all.
    """
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), start=1):
        if not row:
            continue
        m = {}
        for ci, v in enumerate(row):
            key = str(v or "").strip().lower()
            if key in HEADERS and HEADERS[key] not in m:
                m[HEADERS[key]] = ci
        # a real header row names the code plus a few more columns; one stray cell
        # reading "name" somewhere in a title block does not
        if "code" in m and len(m) >= 4:
            return ri, m
    return None, None


def _active(v):
    """'yes'/'no', 1/0, TRUE/FALSE — the exporter writes words, people type anything."""
    s = str(v).strip().lower()
    if s in ("no", "n", "0", "false", "inactive", "ไม่", "ปิด"):
        return 0
    return 1


OTHER_CAT = "Other"


def _norm_name(s):
    return " ".join(str(s or "").split()).lower()


def _other_code(c, factory_id, floor, name):
    """The code for a register row that arrived without one.

    A row with a name and no asset id is not a machine — it is a toilet, a building, a
    length of pipe — and the plant still maintains it, so it belongs in the register
    rather than in a list of things the import threw away. It cannot go in without a
    code, though: the code is what a job, a QR label and a PM schedule all hang off,
    and it is what the next import matches on. So one is minted here, `W01OT07` shaped
    like every other code, and the row is matched **by name** from then on — the same
    sheet imported twice updates that asset instead of creating a second toilet.

    Returns (code, is_new).
    """
    nn = _norm_name(name)
    for r in c.execute("SELECT code,name FROM machines WHERE factory_id=? AND category=?",
                       (factory_id, OTHER_CAT)):
        if _norm_name(r["name"]) == nn:
            return r["code"], False
    fl = str(floor or "").strip() or "1"
    fl = ("0" + fl)[-2:] if fl.isdigit() else "01"
    pre = f"W{fl}OT"
    top = 0
    for r in c.execute("SELECT code FROM machines WHERE factory_id=? AND code LIKE ?",
                       (factory_id, pre + "%")):
        tail = str(r["code"] or "")[len(pre):]
        if tail.isdigit():
            top = max(top, int(tail))
    return f"{pre}{top + 1:02d}", True


def _cell(v):
    if v in (None, ""):
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def import_xlsx(raw_bytes, factory_id, c):
    """Parse workbook bytes and upsert into machines for `factory_id`.

    `c` is a db() connection. Returns {read, inserted, updated, skipped}.
    Raises ValueError with a friendly message on a bad/empty workbook.
    """
    try:
        import openpyxl
    except ImportError:
        raise ValueError("openpyxl is not installed on the server")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValueError("could not read the file — is it a valid .xlsx?")
    if "Review" not in wb.sheetnames:
        # a PM Plan workbook dropped into the asset importer by mistake
        import re as _re
        if any(_re.search(r"[ก-๙]", s) for s in wb.sheetnames) and len(wb.sheetnames) > 20:
            raise ValueError("ไฟล์นี้ดูเหมือนแผน PM ไม่ใช่ทะเบียนสินทรัพย์ — ใช้ปุ่ม “อัปโหลดแผน PM” / "
                             "this looks like the PM plan, not the asset register — "
                             "use the “Upload PM plan” button")
    ws = wb["Review"] if "Review" in wb.sheetnames else wb.active

    # An export-shaped sheet is read by its headers. The Review template keeps the
    # positional reader below, so a workbook that already worked cannot start behaving
    # differently because a header happened to match.
    if "Review" not in wb.sheetnames:
        hrow, hmap = _header_row(ws)
        if hrow:
            return _import_by_header(ws, wb, hrow, hmap, factory_id, c)

    read = inserted = updated = skipped = 0
    for row in ws.iter_rows(min_row=6, values_only=True):     # rows 1-5 = title/header/example
        if row is None or len(row) < 3:
            skipped += 1
            continue
        code = _cell(row[1])
        if not code or code == "W01XX99":                     # empty / example row
            skipped += 1
            continue
        rec = {f: _cell(row[idx]) for f, idx in MAP}
        # scoped to THIS plant: an asset code names one machine inside a plant, and
        # the same code in another plant is a different machine, not this one
        exists = c.execute("SELECT id FROM machines WHERE code=? AND factory_id=?",
                           (code, factory_id)).fetchone()
        if exists:
            upd = FIELDS[1:] + ["factory_id", "active"]        # keep the key, refresh the rest
            c.execute(f"UPDATE machines SET {','.join(k + '=?' for k in upd)} WHERE id=?",
                      (*[rec[f] for f in FIELDS[1:]], factory_id, 1, exists["id"]))
            updated += 1
        else:
            cols = FIELDS + ["factory_id", "active"]
            ph = ",".join("?" * len(cols))
            c.execute(f"INSERT INTO machines({','.join(cols)}) VALUES({ph})",
                      (*[rec[f] for f in FIELDS], factory_id, 1))
            inserted += 1
        read += 1
    wb.close()
    if not read:
        raise ValueError("ไม่พบรายการสินทรัพย์ในไฟล์นี้ (คาดว่าชีต Review รหัสเครื่องอยู่คอลัมน์ B "
                         "เริ่มแถวที่ 6) / no asset rows found — expected a Review sheet with the "
                         "asset code in column B from row 6")
    c.commit()
    return {"read": read, "inserted": inserted, "updated": updated, "skipped": skipped}


def _import_by_header(ws, wb, hrow, hmap, factory_id, c):
    """Upsert an export-shaped sheet, writing only the columns the file actually has.

    A planner exporting the register to fill in one column must not have the other
    fifteen blanked out on the way back, so a column absent from the sheet is left
    alone on the machine, and a cell left empty in a column that IS present is written
    as empty — that is how you clear a value on purpose.
    """
    fields = [f for f in hmap if f != "code"]
    name_ci = hmap.get("name")
    floor_ci = hmap.get("floor")
    read = inserted = updated = skipped = 0
    # Rows left out, and why — a bare count once told a planner that six of their rows
    # had gone somewhere without saying which. Only rows that really cannot be placed
    # land here now; a row with a name is given a code and imported instead.
    dropped = []
    # Rows that arrived with no asset id and were brought in as Other, so the planner
    # can see exactly which codes were minted for them and rename or recategorise later.
    noid = []
    for rn, row in enumerate(ws.iter_rows(min_row=hrow + 1, values_only=True), start=hrow + 1):
        if not row or not any(v not in (None, "") for v in row):
            skipped += 1
            continue                                   # a truly empty row is not news
        code = _cell(row[hmap["code"]]) if hmap["code"] < len(row) else ""
        if code == "W01XX99":                          # the template's example row
            skipped += 1
            continue
        force_other = False
        if not code:
            nm = _cell(row[name_ci]) if (name_ci is not None and name_ci < len(row)) else ""
            if not nm:
                skipped += 1
                if len(dropped) < 25:
                    dropped.append({"row": rn, "name": "", "why": "no asset id and no name"})
                continue
            fl = _cell(row[floor_ci]) if (floor_ci is not None and floor_ci < len(row)) else ""
            code, is_new = _other_code(c, factory_id, fl, nm)
            force_other = True
            if len(noid) < 40:
                noid.append({"row": rn, "name": nm, "code": code, "new": is_new})
        rec = {}
        for f in fields:
            ci = hmap[f]
            v = row[ci] if ci < len(row) else ""
            if f == "active":
                rec[f] = _active(v)
            elif f == "pm_freq_days":
                try:
                    rec[f] = int(float(_cell(v) or 0))
                except ValueError:
                    rec[f] = 0
            else:
                rec[f] = _cell(v)
        # A row with no asset id is not a machine — it is a structure, a utility, a
        # building. Filed as Other until somebody decides what it really is, and the
        # category is forced rather than defaulted so a blank Category cell in the sheet
        # cannot quietly wipe it on the next import.
        if force_other:
            rec["category"] = OTHER_CAT
        exists = c.execute("SELECT id FROM machines WHERE code=? AND factory_id=?",
                           (code, factory_id)).fetchone()
        if exists:
            cols = list(rec) + ["factory_id"]
            c.execute(f"UPDATE machines SET {','.join(k + '=?' for k in cols)} WHERE id=?",
                      (*[rec[f] for f in rec], factory_id, exists["id"]))
            updated += 1
        else:
            rec.setdefault("active", 1)
            cols = ["code"] + list(rec) + ["factory_id"]
            ph = ",".join("?" * len(cols))
            c.execute(f"INSERT INTO machines({','.join(cols)}) VALUES({ph})",
                      (code, *[rec[f] for f in rec], factory_id))
            inserted += 1
        read += 1
    wb.close()
    if not read:
        raise ValueError("ไม่พบรายการสินทรัพย์ในไฟล์นี้ / no asset rows found under the header row")
    c.commit()
    return {"read": read, "inserted": inserted, "updated": updated, "skipped": skipped,
            "dropped": dropped, "noid": noid}
